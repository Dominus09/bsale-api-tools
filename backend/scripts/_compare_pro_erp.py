"""Comparación visual PRO vs ERP (datos de prueba controlados)."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from backend.services.analytics_v2.commercial_premium_excel import (
    commercial_premium_review_xlsx,
    commercial_premium_xlsx,
)
from backend.tests.test_commercial_premium_excel import _controlled_payload

OUT = Path(__file__).resolve().parent / "_fase25_xlsx_out"
PRO_REV = Path(r"c:\Users\user\OneDrive\MIO\Escritorio\La_Quillotana_Excel_Revision_PRO.xlsx")
PRO_FULL = Path(r"c:\Users\user\OneDrive\MIO\Escritorio\La_Quillotana_Auditoria_Completa_PRO.xlsx")


def _fill(cell) -> str | None:
    try:
        return cell.fill.fgColor.rgb
    except Exception:
        return None


def compare(label: str, erp_bytes: bytes, pro_path: Path) -> None:
    erp = load_workbook(BytesIO(erp_bytes))
    pro = load_workbook(pro_path, read_only=False, data_only=False)
    print("===", label, "===")
    print("sheets ERP", erp.sheetnames)
    print("sheets PRO", pro.sheetnames)
    for name in erp.sheetnames:
        if name not in pro.sheetnames:
            print("  missing in PRO", name)
            continue
        we, wp = erp[name], pro[name]
        a1e, a1p = we["A1"].value, wp["A1"].value
        print(f" [{name}]")
        print(f"  freeze ERP={we.freeze_panes} PRO={wp.freeze_panes}")
        print(f"  filter ERP={we.auto_filter.ref} PRO={wp.auto_filter.ref}")
        print(f"  A1 ERP={a1e!r}")
        print(f"  A1 PRO={a1p!r}")
        print(f"  A1 fill ERP={_fill(we['A1'])} PRO={_fill(wp['A1'])}")

    if "Críticos y anomalías" in erp.sheetnames:
        we = erp["Críticos y anomalías"]
        wp = pro["Críticos y anomalías"]
        he = [we.cell(5, c).value for c in range(1, 14)]
        hp = [wp.cell(5, c).value for c in range(1, 14)]
        print(" anomalies hdr1-13 ERP", he)
        print(" anomalies hdr1-13 PRO", hp)
        he_all = [we.cell(5, c).value for c in range(1, we.max_column + 1)]
        hp_all = [wp.cell(5, c).value for c in range(1, wp.max_column + 1)]
        print(
            " prices ERP",
            [h for h in he_all if h and str(h).lower().startswith("precio")],
        )
        print(
            " prices PRO",
            [h for h in hp_all if h and str(h).upper().startswith("PRECIO")],
        )
        print(" office row6 ERP", we.cell(6, 13).value, "PRO", wp.cell(6, 13).value)

    if "Por revisar" in erp.sheetnames:
        he = [erp["Por revisar"].cell(8, c).value for c in range(1, 12)]
        hp = [pro["Por revisar"].cell(8, c).value for c in range(1, 12)]
        print(" review hdr ERP", he)
        print(" review hdr PRO", hp)
    pro.close()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = _controlled_payload()
    full = commercial_premium_xlsx(payload)
    rev = commercial_premium_review_xlsx(payload)
    (OUT / "ERP_Auditoria_Completa_PRUEBA.xlsx").write_bytes(full)
    (OUT / "ERP_Excel_Revision_PRUEBA.xlsx").write_bytes(rev)
    print("wrote", OUT, "sizes", len(full), len(rev))
    compare("REVISION", rev, PRO_REV)
    compare("FULL", full, PRO_FULL)


if __name__ == "__main__":
    main()
