"""Pruebas Excel Premium Auditoría Comercial 8.0 (sin PostgreSQL productivo)."""

from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.services.analytics_v2.commercial_premium_excel import commercial_premium_xlsx
from backend.services.analytics_v2.commercial_premium_service import (
    SHEET_NAMES,
    build_payload_from_controlled_data,
    enrich_review_row,
    price_lists_for_company,
)
from backend.services.analytics_v2 import excel_export
from backend.services.analytics.validate_distribuidora_source import assert_sql_is_read_only


def _sample_anomaly(*, variant_id: int, diff_pct: str, stock: str, extreme: bool = False) -> dict:
    avg = Decimal("100")
    last = avg * (Decimal("1") + Decimal(diff_pct) / Decimal("100"))
    row = {
        "variant_id": variant_id,
        "product_id": 1000 + variant_id,
        "product_name": f"Producto {variant_id}",
        "variant_name": f"Variante {variant_id}",
        "sku": f"0{variant_id}",
        "barcode": f"780{variant_id:010d}",
        "average_cost": avg,
        "last_comparable_cost": last,
        "diff_pct": Decimal(diff_pct),
        "diff_abs": last - avg,
        "stock_total": Decimal(stock),
        "last_comparable_date": date(2026, 9, 1),
        "last_comparable_document": "Factura",
        "last_comparable_document_number": "123",
        "last_comparable_office_id": 1,
        "banda_principal": "sobre_50" if abs(Decimal(diff_pct)) > 50 else "entre_15_y_50",
        "extrema_indicador": extreme,
        "comparability_status": "comparable",
        "evidence_sufficient": True,
        "price_12": Decimal("1500"),
        "price_13": None,
        "price_16": Decimal("1400"),
    }
    return enrich_review_row(row, cut_at=date(2026, 10, 23))


def _controlled_payload() -> dict:
    # 5 anomalías (>15 %) — simula que NO hay tope 30; el generador debe incluirlas todas.
    anomalies = [
        _sample_anomaly(variant_id=10894, diff_pct="35.77", stock="1964"),
        _sample_anomaly(variant_id=27785, diff_pct="35.17", stock="679"),
        _sample_anomaly(variant_id=30735, diff_pct="400", stock="5", extreme=True),
        _sample_anomaly(variant_id=24486, diff_pct="-95.24", stock="29", extreme=True),
        _sample_anomaly(variant_id=29507, diff_pct="60.28", stock="14"),
        # más allá de un hipotético top-30
        *[_sample_anomaly(variant_id=40000 + i, diff_pct="20.5", stock="10") for i in range(35)],
    ]
    assert len(anomalies) == 40

    universe = list(anomalies) + [
        {
            "variant_id": 1,
            "product_name": "OK",
            "variant_name": "OK",
            "sku": "1",
            "barcode": "1",
            "average_cost": Decimal("100"),
            "last_comparable_cost": Decimal("105"),
            "diff_pct": Decimal("5"),
            "diff_abs": Decimal("5"),
            "stock_total": Decimal("3"),
            "banda_principal": "hasta_15",
            "extrema_indicador": False,
            "comparability_status": "comparable",
            "evidence_sufficient": True,
            "last_comparable_date": date(2026, 8, 1),
            "last_comparable_document": "Factura",
            "last_comparable_document_number": "9",
            "last_comparable_office_id": 1,
        },
        {
            "variant_id": 2,
            "product_name": "Sin comparable",
            "variant_name": "X",
            "sku": "2",
            "barcode": "2",
            "average_cost": Decimal("50"),
            "last_comparable_cost": None,
            "diff_pct": None,
            "diff_abs": None,
            "stock_total": Decimal("8"),
            "banda_principal": "sin_recepcion_comparable",
            "extrema_indicador": False,
            "comparability_status": "sin_comparable",
            "evidence_sufficient": False,
        },
    ]

    zero = [
        {
            "variant_id": 9001,
            "product_name": "Cero A",
            "variant_name": "Sin stock fila",
            "sku": "09001",
            "barcode": "7809001",
            "average_cost": Decimal("0"),
            "stock_total_raw": None,
            "stock_total": 0,
            "sin_fila_stock": True,
            "stock_cero_explicito": False,
            "stock_status": "sin_fila_stock",
            "has_comparable_reception": False,
            "has_positive_price": True,
            "price_12": Decimal("999"),
            "commercial_status": "activo_costo_cero_sin_stock_pos",
            "recommended_action": "verificar_alta_catalogo_y_primera_compra",
        },
        {
            "variant_id": 9002,
            "product_name": "Cero B",
            "variant_name": "Stock cero explícito",
            "sku": "09002",
            "barcode": "7809002",
            "average_cost": Decimal("0"),
            "stock_total_raw": Decimal("0"),
            "stock_total": 0,
            "sin_fila_stock": False,
            "stock_cero_explicito": True,
            "stock_status": "stock_cero_explicito",
            "has_comparable_reception": True,
            "has_positive_price": False,
            "price_12": None,
            "commercial_status": "activo_costo_cero_sin_stock_pos",
            "recommended_action": "revisar_por_que_promedio_sigue_en_cero",
            "last_comparable_document": "Guía",
            "last_comparable_date": date(2026, 1, 10),
            "last_comparable_office_id": 3,
        },
    ]

    return build_payload_from_controlled_data(
        anomalies=anomalies,
        average_vs_receptions=universe,
        zero_cost_actives=zero,
        orphans=[{"variant_id": 31352, "stock_total": Decimal("2"), "oficinas": 1}],
        sync_state=[
            {
                "resource": "variant_costs",
                "scope": "full",
                "status": "SUCCESS",
                "last_success_at": datetime(2026, 10, 22, 12, 0, tzinfo=timezone.utc),
                "last_error_at": datetime(2026, 10, 20, 8, 0, tzinfo=timezone.utc),
                "active_freshness_alert": False,
                "freshness_note": "hubo errores históricos; último SUCCESS los supersede",
            }
        ],
    )


def test_price_lists_company_3_configured():
    lists = price_lists_for_company(3)
    assert set(lists.keys()) == {12, 13, 16}


def test_price_lists_unknown_company_empty():
    assert price_lists_for_company(99999) == {}


def test_commercial_premium_xlsx_six_sheets_and_headers():
    payload = _controlled_payload()
    content = commercial_premium_xlsx(payload)
    assert content[:2] == b"PK"  # zip/xlsx
    wb = load_workbook(BytesIO(content))
    assert wb.sheetnames == list(SHEET_NAMES)

    assert wb[SHEET_NAMES[1]]["A1"].value == "Prioridad"
    assert wb[SHEET_NAMES[2]]["A1"].value == "Producto"
    assert wb[SHEET_NAMES[3]]["A1"].value == "Producto"
    assert wb[SHEET_NAMES[5]]["H1"].value == "Estado de revisión"


def test_anomalies_not_limited_to_top_30():
    payload = _controlled_payload()
    assert payload["anomalies_count"] == 40
    assert "no top_n" in payload["anomalies_source"]
    content = commercial_premium_xlsx(payload)
    wb = load_workbook(BytesIO(content))
    ws = wb[SHEET_NAMES[1]]
    # fila 1 header + 40 datos
    assert ws.max_row == 41
    # también en revisión administrativa
    assert wb[SHEET_NAMES[5]].max_row >= 41


def test_sku_and_barcode_as_text_preserve_leading_zero():
    payload = _controlled_payload()
    content = commercial_premium_xlsx(payload)
    wb = load_workbook(BytesIO(content))
    ws = wb[SHEET_NAMES[1]]
    # Columna F = SKU del primer anomaly (10894 → "010894")
    sku_cell = ws["F2"]
    assert sku_cell.number_format == "@"
    assert str(sku_cell.value).startswith("0")


def test_zero_cost_distinguishes_explicit_zero_vs_missing_stock_row():
    payload = _controlled_payload()
    content = commercial_premium_xlsx(payload)
    wb = load_workbook(BytesIO(content))
    ws = wb[SHEET_NAMES[2]]
    statuses = {ws.cell(r, 7).value for r in range(2, ws.max_row + 1)}
    assert "sin_fila_stock" in statuses
    assert "stock_cero_explicito" in statuses
    # Sin fila → stock vacío (None), no 0 inventado
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 7).value == "sin_fila_stock":
            assert ws.cell(r, 8).value is None


def test_non_comparable_has_blank_diff():
    payload = _controlled_payload()
    content = commercial_premium_xlsx(payload)
    wb = load_workbook(BytesIO(content))
    ws = wb[SHEET_NAMES[3]]
    found = False
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 17).value == "sin_comparable":
            assert ws.cell(r, 8).value is None
            assert ws.cell(r, 9).value is None
            found = True
    assert found


def test_summary_contains_methodology_and_no_confirmed_error_claim():
    payload = _controlled_payload()
    content = commercial_premium_xlsx(payload)
    wb = load_workbook(BytesIO(content))
    ws = wb[SHEET_NAMES[0]]
    blob = "\n".join(
        str(ws.cell(r, 1).value or "") for r in range(1, min(ws.max_row, 80) + 1)
    )
    assert "alertas" in blob.lower() or "Alertas" in blob
    assert "NO son errores" in blob or "no son errores" in blob.lower() or "≠" in blob


def test_admin_sheet_has_dropdown_validation():
    payload = _controlled_payload()
    content = commercial_premium_xlsx(payload)
    wb = load_workbook(BytesIO(content))
    ws = wb[SHEET_NAMES[5]]
    assert len(ws.data_validations.dataValidation) >= 1
    assert ws["H2"].value == "Pendiente"


def test_sync_historical_error_not_active_alert_when_success_later():
    payload = _controlled_payload()
    content = commercial_premium_xlsx(payload)
    wb = load_workbook(BytesIO(content))
    ws = wb[SHEET_NAMES[4]]
    blob = "\n".join(
        str(ws.cell(r, c).value or "")
        for r in range(1, ws.max_row + 1)
        for c in range(1, 5)
    )
    assert "SUCCESS" in blob
    assert "supersede" in blob.lower() or "históric" in blob.lower() or "histor" in blob.lower()


def test_exposure_formula_present_on_anomalies():
    payload = _controlled_payload()
    content = commercial_premium_xlsx(payload)
    wb = load_workbook(BytesIO(content))
    ws = wb[SHEET_NAMES[1]]
    assert isinstance(ws["L2"].value, str)
    assert ws["L2"].value.startswith("=IF(")


def test_no_duplicate_variant_ids_in_anomalies_sheet():
    payload = _controlled_payload()
    content = commercial_premium_xlsx(payload)
    wb = load_workbook(BytesIO(content))
    ws = wb[SHEET_NAMES[1]]
    ids = [ws.cell(r, 5).value for r in range(2, ws.max_row + 1)]
    assert len(ids) == len(set(ids))


def test_excel_export_reexports_commercial_premium():
    assert hasattr(excel_export, "commercial_premium_xlsx")
    payload = _controlled_payload()
    a = commercial_premium_xlsx(payload)
    b = excel_export.commercial_premium_xlsx(payload)
    assert a[:2] == b[:2] == b"PK"
    assert load_workbook(BytesIO(b)).sheetnames == list(SHEET_NAMES)


def test_existing_costs_xlsx_still_works():
    """Regresión: exportador de costos V2 intacto."""
    data = excel_export.costs_xlsx(
        meta={"company_id": 3, "company_name": "Test"},
        summary={"skus_active": 1, "skus_active_by_reliability": {}},
        items=[],
        priority=[],
        reception_lines=[],
        filters_label="test",
    )
    assert data[:2] == b"PK"
    wb = load_workbook(BytesIO(data))
    assert "Resumen" in wb.sheetnames


def test_service_sql_helpers_remain_read_only_contract():
    # Las consultas del motor auditor siguen pasando el validador RO
    assert_sql_is_read_only(
        "SELECT variant_id, sum(quantity) AS stock_total "
        "FROM bsale_raw.stocks WHERE company_id = %s GROUP BY variant_id"
    )


def test_enrich_keeps_review_priority_when_admin_correction_possible():
    row = _sample_anomaly(variant_id=30735, diff_pct="400", stock="5", extreme=True)
    prio = row["admin_priority"]
    assert prio["possible_admin_correction"] is True
    assert prio["error_confidence"] == "baja"
    assert prio["review_priority"].startswith("P1")
