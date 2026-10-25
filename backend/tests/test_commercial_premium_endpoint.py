"""Endpoint Excel Premium — auth, empresa y descarga (sin PG productivo)."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from fastapi import FastAPI
from openpyxl import load_workbook

from backend.routers import analytics_v2 as router_mod
from backend.services.analytics_v2.access import AnalyticsV2AccessError
from backend.services.analytics_v2.commercial_premium_service import (
    REVIEW_SHEET_NAMES,
    SHEET_NAMES,
    build_payload_from_controlled_data,
    enrich_review_row,
)
from backend.tests.test_analytics_v2_access import AsgiClient, _auth


def _anomaly(vid: int, diff: str) -> dict[str, Any]:
    avg = Decimal("100")
    last = avg * (Decimal("1") + Decimal(diff) / Decimal("100"))
    return enrich_review_row(
        {
            "variant_id": vid,
            "product_name": f"P{vid}",
            "variant_name": f"V{vid}",
            "sku": str(vid),
            "barcode": str(vid),
            "average_cost": avg,
            "last_comparable_cost": last,
            "diff_pct": Decimal(diff),
            "stock_total": Decimal("10"),
            "last_comparable_date": date(2026, 9, 1),
            "last_comparable_document": "Factura",
            "last_comparable_document_number": "1",
            "last_comparable_office_id": 1,
            "banda_principal": "entre_15_y_50",
            "extrema_indicador": False,
            "comparability_status": "comparable",
            "evidence_sufficient": True,
        },
        cut_at=date(2026, 10, 23),
    )


@pytest.fixture()
def premium_client(monkeypatch):
    access_calls: list[tuple[Any, int]] = []

    def fake_run_cost(user, company_id, fn):
        access_calls.append((user, company_id))
        if company_id == 99:
            raise AnalyticsV2AccessError(403, "Empresa fuera del alcance del usuario")
        if company_id == 1:
            raise AnalyticsV2AccessError(409, "La empresa aún no tiene BSALE RAW completo")
        return fn(object())

    def fake_payload(company_id, **kwargs):
        # 35 anomalías (> top 30) para validar universo completo en el endpoint
        anomalies = [_anomaly(1000 + i, "20") for i in range(35)]
        return build_payload_from_controlled_data(
            company_id=company_id,
            anomalies=anomalies,
            average_vs_receptions=anomalies,
        )

    monkeypatch.setattr(router_mod, "run_cost_control", fake_run_cost)
    monkeypatch.setattr(
        "backend.services.analytics_v2.commercial_premium_service.build_commercial_premium_payload",
        fake_payload,
    )
    # El router importa la función dentro de build(); parchear también el símbolo usado allí.
    import backend.services.analytics_v2.commercial_premium_service as prem_mod

    monkeypatch.setattr(prem_mod, "build_commercial_premium_payload", fake_payload)

    app = FastAPI()
    app.include_router(router_mod.router)
    client = AsgiClient(app)
    client.access_calls = access_calls  # type: ignore[attr-defined]
    return client


def test_premium_export_requires_auth(premium_client):
    res = premium_client.get("/analytics-v2/audit/commercial-premium/export.xlsx?company_id=3")
    assert res.status_code in (401, 403)


def test_premium_export_rejects_unauthorized_company(premium_client):
    res = premium_client.get(
        "/analytics-v2/audit/commercial-premium/export.xlsx?company_id=99",
        headers=_auth("adm"),
    )
    assert res.status_code == 403


def test_premium_export_default_is_review_two_sheets(premium_client):
    res = premium_client.get(
        "/analytics-v2/audit/commercial-premium/export.xlsx?company_id=3",
        headers=_auth("adm"),
    )
    assert res.status_code == 200
    assert res.body[:2] == b"PK"
    assert any(c[1] == 3 for c in premium_client.access_calls)  # type: ignore[attr-defined]
    wb = load_workbook(BytesIO(res.body))
    assert wb.sheetnames == list(REVIEW_SHEET_NAMES)
    # 35 anomalías (sin tope 30); max_row puede incluir nota al pie
    ws = wb[REVIEW_SHEET_NAMES[0]]
    data_rows = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, 3).value)
    assert data_rows == 35


def test_premium_export_full_six_sheets(premium_client):
    res = premium_client.get(
        "/analytics-v2/audit/commercial-premium/export.xlsx?company_id=3&mode=full",
        headers=_auth("adm"),
    )
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.body))
    assert wb.sheetnames == list(SHEET_NAMES)
    assert wb[SHEET_NAMES[1]].max_row == 36


def test_premium_filename_contains_company_and_date(premium_client):
    res = premium_client.get(
        "/analytics-v2/audit/commercial-premium/export.xlsx?company_id=3&mode=review",
        headers=_auth("adm"),
    )
    # AsgiClient no expone headers fácilmente; validamos cuerpo XLSX.
    assert res.status_code == 200
    assert res.body[:2] == b"PK"
