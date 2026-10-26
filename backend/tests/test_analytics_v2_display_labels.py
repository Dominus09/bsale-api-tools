"""Etiquetas comerciales administrativas (listas y sucursales)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from backend.services.analytics_v2.commercial_premium_excel import (
    commercial_premium_review_xlsx,
    commercial_premium_xlsx,
)
from backend.services.analytics_v2.commercial_premium_service import (
    REVIEW_SHEET_NAMES,
    SHEET_NAMES,
    build_payload_from_controlled_data,
    price_lists_for_company,
)
from backend.services.analytics_v2.commercial_scope import (
    evaluated_price_list_names,
    resolve_price_list_scope,
)
from backend.services.analytics_v2.contracts import PriceListScopeStatus
from backend.services.analytics_v2.display_labels import (
    office_visible_name,
    price_column_header,
    price_list_visible_name,
)


def test_price_list_visible_prefers_commercial_never_lista_id():
    assert price_list_visible_name(
        name="SUPERMERCADO LA QUILLOTANA", commercial_label="Supermercado"
    ) == "Supermercado"
    assert price_list_visible_name(name="MELINKA", commercial_label="Melinka") == "Melinka"
    assert "12" not in price_list_visible_name(name="Melinka", price_list_id=12)
    assert "Lista 12" not in price_list_visible_name(
        name=None, commercial_label=None, price_list_id=12
    )
    assert "requiere corrección" in price_list_visible_name(name=None, price_list_id=12)


def test_price_column_header_without_id():
    assert price_column_header("Supermercado") == "Precio Supermercado"
    assert "(#" not in price_column_header("Ruta Facturas")
    assert price_column_header("Ruta Facturas") == "Precio Ruta Facturas"


def test_office_visible_unifies_quillotana_roman():
    assert office_visible_name(name="QUILLOTANA I") == "Quillotana I"
    assert office_visible_name(name="Quillotana 1 / Oficina 4") == "Quillotana I"
    assert office_visible_name(name="Quillotana 2") == "Quillotana II"
    assert office_visible_name(name=None, office_id=4, company_id=3) == "Quillotana I"
    assert office_visible_name(name=None, office_id=1, company_id=3) == "Bodega Central"
    assert "Oficina 4" not in office_visible_name(name=None, office_id=4)
    assert "requiere corrección" in office_visible_name(name=None, office_id=4)


def test_evaluated_lists_exclude_inactive_and_other_company_ids():
    rows = [
        {"price_list_id": 12, "name": "SUPERMERCADO LA QUILLOTANA", "state": 0, "missing_since": None},
        {"price_list_id": 13, "name": "RUTA/WEB (BOLETA O FACTURA)", "state": 0, "missing_since": None},
        {"price_list_id": 16, "name": "MELINKA", "state": 0, "missing_since": None},
        {"price_list_id": 14, "name": "COMODITI ALCOHOLES", "state": 1, "missing_since": None},
        {"price_list_id": 99, "name": "OTRA", "state": 0, "missing_since": None},
    ]
    scope = resolve_price_list_scope(3, rows)
    names = evaluated_price_list_names(scope)
    assert set(names) == {12, 13, 16}
    assert 14 not in names
    assert 99 not in names
    assert names[13] == "Ruta Facturas"
    assert all("(#" not in v and "Lista " not in v for v in names.values())
    inactive = next(e for e in scope.entries if e.price_list_id == 14)
    assert inactive.status is not PriceListScopeStatus.EVALUADA


def test_price_lists_for_company_with_raw_filters_active():
    rows = [
        {"price_list_id": 12, "name": "SUPERMERCADO LA QUILLOTANA", "state": 0, "missing_since": None},
        {"price_list_id": 14, "name": "COMODITI", "state": 1, "missing_since": None},
    ]
    out = price_lists_for_company(3, rows)
    assert 12 in out and 14 not in out
    assert out[12] == "Supermercado"


def test_excel_headers_use_commercial_names_not_ids():
    payload = build_payload_from_controlled_data(
        price_lists={12: "Supermercado La Quillotana", 13: "Ruta Facturas", 16: "Melinka"},
        anomalies=[
            {
                "variant_id": 1,
                "product_type_name": "CERVEZA",
                "product_name": "CRISTAL",
                "variant_name": "LATA 470",
                "sku": "01",
                "barcode": "7801",
                "average_cost": 100,
                "last_comparable_cost": 120,
                "diff_pct": 20,
                "stock_total": 5,
                "last_comparable_office_id": 4,
                "admin_priority": {"review_priority": "P2_importante", "priority_score": 50},
                "price_12": 1500,
                "price_13": None,
                "price_16": 1400,
            }
        ],
    )
    content = commercial_premium_xlsx(payload)
    wb = load_workbook(BytesIO(content))
    assert wb.sheetnames == list(SHEET_NAMES)
    ws = wb[SHEET_NAMES[1]]
    # Banner filas 1–3; encabezados en fila 5
    headers = [ws.cell(5, c).value for c in range(1, ws.max_column + 1)]
    assert any(h == "Precio Supermercado La Quillotana" for h in headers)
    assert any(h == "Precio Melinka" for h in headers)
    assert any(h == "Precio Ruta Facturas" for h in headers)
    assert not any(h and "(#" in str(h) for h in headers)
    assert not any(h and str(h).lower().startswith("precio lista") for h in headers)
    assert not any(h and str(h).upper() in {"PRECIO 12", "PRECIO 13", "PRECIO 16"} for h in headers)
    office_col = headers.index("OFICINA") + 1
    assert ws.cell(6, office_col).value == "Quillotana I"
    price_cols = [i for i, h in enumerate(headers, start=1) if h and str(h).startswith("Precio ")]
    assert any(ws.cell(6, c).value == "Sin precio" for c in price_cols)

    rev = commercial_premium_review_xlsx(payload)
    wb2 = load_workbook(BytesIO(rev))
    assert wb2.sheetnames == list(REVIEW_SHEET_NAMES)
