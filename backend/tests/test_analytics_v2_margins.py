"""Márgenes V2: fórmulas Decimal, motor de estados, excepciones promocionales y servicio."""

from __future__ import annotations

import re
import time
from datetime import date, datetime, timezone
from decimal import Decimal
from typing import Any

import pytest

from backend.services.analytics.validate_distribuidora_source import assert_sql_is_read_only
from backend.services.analytics_v2 import cost_control_service as cost_service_mod
from backend.services.analytics_v2 import margins_repo as repo_mod
from backend.services.analytics_v2 import margins_service as margins_service_mod
from backend.services.analytics_v2.access import CompanyScope
from backend.services.analytics_v2.contracts import (
    CostBasis,
    CostReliability,
    MarginComplianceStatus as S,
    PolicyReviewStatus,
    PromotionCheck,
    RuleIssue,
    TaxProfileStatus,
)
from backend.services.analytics_v2.cost_control_service import CostControlV2Service
from backend.services.analytics_v2.margin_engine import (
    CostFacts,
    PriceFact,
    RuleFact,
    evaluate,
    price_issue,
    rule_issues,
)
from backend.services.analytics_v2.margin_math import (
    ceiling_price,
    margin_on_sale_pct,
    markup_pct,
    round_pct,
    target_price,
)
from backend.services.analytics_v2.margins_repo import MarginsV2Repository
from backend.services.analytics_v2.margins_service import MarginsV2Service, MarginsValidationError
from backend.services.analytics_v2.promotion_exceptions import (
    NOT_INTEGRATED,
    ApprovalStatus,
    ExceptionKind,
    ExceptionResolution,
    PromotionException,
    resolve_exception,
)
from backend.services.analytics_v2.tax_profile import TaxEntry, TaxProfile
from backend.tests.test_analytics_v2_cost_control import FakeRepo, fact, purchase

D = Decimal
TODAY = date(2026, 10, 13)
SCOPE = CompanyScope(company_id=3, company_name="La Quillotana SpA")
IVA = TaxProfile(
    status=TaxProfileStatus.IVA,
    reason="IVA 19%",
    taxes=(TaxEntry(1, "IVA 19%", D("19"), "iva"),),
    total_rate_pct=D("19"),
)
ANTICIPADO = TaxProfile(status=TaxProfileStatus.IVA_ANTICIPADO, reason="IVA HARI.")


# ---------------------------------------------------------------- fórmulas


def test_markup_and_margin_on_sale_use_exact_decimal():
    assert markup_pct(D("1190"), D("1000")) == D("19")
    assert margin_on_sale_pct(D("1250"), D("1000")) == D("20")
    assert round_pct(markup_pct(D("1000"), D("3"))) == D("33233.33")
    assert isinstance(markup_pct(D("10"), D("3")), Decimal)


def test_formulas_never_divide_by_zero():
    assert markup_pct(D("100"), D("0")) is None
    assert markup_pct(D("100"), None) is None
    assert markup_pct(None, D("10")) is None
    assert margin_on_sale_pct(D("0"), D("10")) is None
    assert margin_on_sale_pct(D("100"), None) is None


def test_target_prices_round_to_peso_on_the_safe_side():
    assert target_price(D("1000.40"), D("20")) == D("1201")
    assert markup_pct(target_price(D("1000.40"), D("20")), D("1000.40")) >= D("20")
    assert ceiling_price(D("1000.40"), D("30")) == D("1300")
    assert ceiling_price(D("1000"), D("0")) is None
    assert target_price(D("0"), D("20")) is None


# ---------------------------------------------------------------- motor


def rule(min_m="20", max_m="30", **kw) -> RuleFact:
    min_d = D(min_m) if min_m is not None else None
    max_d = D(max_m) if max_m is not None else None
    active = kw.pop("active", True)
    return RuleFact(
        rule_id=kw.pop("rule_id", 1),
        price_list_id=kw.pop("price_list_id", 12),
        product_type_id=kw.pop("product_type_id", 1),
        min_markup=min_d,
        max_markup=max_d,
        active=active,
        issues=rule_issues(min_d, max_d, active),
        **kw,
    )


def cost(net="1000", avg="900", reliability=CostReliability.CONSISTENTE, profile=IVA) -> CostFacts:
    return CostFacts(
        average_net=D(avg) if avg is not None else None,
        last_comparable_net=D(net) if net is not None else None,
        reliability=reliability,
        tax_profile=profile,
    )


def price(gross="1500", net=None, missing=False) -> PriceFact:
    return PriceFact(gross=D(gross) if gross is not None else None, net=D(net) if net else None, missing=missing)


def run(**kw) -> dict[str, Any]:
    args = dict(
        price=price(),
        cost=cost(),
        rule=rule(),
        rule_match="tipo_producto",
        cost_basis=CostBasis.ULTIMO_COMPARABLE,
        promotion=NOT_INTEGRATED,
    )
    args.update(kw)
    return evaluate(**args)


def test_within_policy_uses_gross_cost_and_reports_both_formulas():
    r = run(price=price("1500"))
    # costo bruto preliminar = 1000 × 1,19 = 1190
    assert r["costs"]["reference"]["gross"] == "1190.00"
    assert r["markup_pct"] == "26.05"
    assert r["margin_on_sale_pct"] == "20.67"
    assert r["status"] == S.DENTRO_POLITICA
    assert r["gap_to_min_pct"] == "6.05"
    assert r["target_price"] == {"min": "1428", "max": "1547", "provisional": True}


def test_statuses_are_calculated_not_confirmed_while_cost_and_policy_unverified():
    r = run()
    assert r["confirmation"] == "calculado"
    assert set(r["qualifiers"]) == {S.COSTO_NO_VERIFICADO, S.POLITICA_PROVISIONAL}
    assert r["validity"] == {"status": "provisional", "reasons": ["politica_pendiente_revision", "costo_no_verificado"]}


def test_confirmed_only_with_validated_policy_and_verified_cost():
    validated = rule(review_status=PolicyReviewStatus.VALIDADA)
    r = run(rule=validated, net_cost_verified=True, tax_config_verified=True)
    assert r["qualifiers"] == [] and r["confirmation"] == "confirmado"
    assert r["validity"]["status"] == "validado" and r["target_price"]["provisional"] is False


def test_unverified_tax_configuration_keeps_margin_provisional():
    validated = rule(review_status=PolicyReviewStatus.VALIDADA)
    r = run(rule=validated, net_cost_verified=True, tax_config_verified=False)
    assert r["markup_pct"] is not None
    assert S.COSTO_NO_VERIFICADO in r["qualifiers"]
    assert r["validity"]["status"] == "provisional"


def test_below_and_above_policy():
    assert run(price=price("1300"))["status"] == S.BAJO_POLITICA
    assert run(price=price("1600"))["status"] == S.SOBRE_POLITICA


def test_below_policy_is_pending_promotion_check_while_promotions_not_integrated():
    r = run(price=price("1300"))
    assert S.PROMOCION_PENDIENTE_VERIFICACION in r["qualifiers"]
    assert r["promotion"]["check"] == PromotionCheck.FUENTE_NO_INTEGRADA
    assert S.PROMOCION_PENDIENTE_VERIFICACION not in run(price=price("1500"))["qualifiers"]


@pytest.mark.parametrize(
    "p, issue",
    [
        (None, "sin_precio_en_lista"),
        (PriceFact(gross=None, net=None), "precio_bruto_nulo"),
        (PriceFact(gross=D("0"), net=D("0")), "precio_cero_o_negativo"),
        (PriceFact(gross=D("1"), net=D("1")), "precio_simbolico"),
        (PriceFact(gross=D("1500"), net=None, missing=True), "precio_retirado_en_bsale"),
    ],
)
def test_null_zero_symbolic_or_retired_prices_are_sin_precio_without_percentages(p, issue):
    r = run(price=p)
    assert r["status"] == S.SIN_PRECIO and r["reasons"] == [issue]
    assert r["markup_pct"] is None and r["margin_on_sale_pct"] is None
    assert price_issue(p) == issue


@pytest.mark.parametrize("net", [None, "0"])
def test_zero_or_null_cost_is_sin_costo(net):
    r = run(cost=cost(net=net))
    assert r["status"] == S.SIN_COSTO and r["markup_pct"] is None
    assert r["target_price"]["min"] is None


def test_cost_basis_is_explicit_and_never_substituted():
    no_purchase = cost(net=None, avg="900")
    assert run(cost=no_purchase)["status"] == S.SIN_COSTO
    r = run(cost=no_purchase, cost_basis=CostBasis.PROMEDIO_BSALE)
    assert r["costs"]["reference"]["basis"] == "promedio_bsale"
    assert r["costs"]["reference"]["net"] == "900"
    assert r["costs"]["bsale_average_net"] == "900" and r["costs"]["last_reception_net"] is None


def test_missing_or_inactive_rule_is_never_compliance():
    r = run(rule=None)
    assert r["status"] == S.SIN_REGLA and r["markup_pct"] is None and r["rule"] is None
    inactive = run(rule=rule(active=False))
    assert inactive["status"] == S.SIN_REGLA and inactive["reasons"] == ["regla_no_aplicable"]
    assert RuleIssue.REGLA_INACTIVA in inactive["rule"]["issues"]


def test_max_zero_is_flagged_and_never_produces_sobre_politica():
    r = run(rule=rule("0", "0"), price=price("5000"))
    assert r["status"] == S.DENTRO_POLITICA
    assert r["rule"]["issues"] == [RuleIssue.MAXIMO_CERO]
    assert r["target_price"]["max"] is None


def test_rule_issue_detection():
    assert rule_issues(D("30"), D("20"), True) == (RuleIssue.MINIMO_MAYOR_QUE_MAXIMO,)
    assert rule_issues(None, None, None) == (RuleIssue.MINIMO_NULO,)
    assert rule_issues(D("10"), None, True) == ()


def test_pending_tax_profile_requires_tax_review_without_percentages():
    r = run(cost=cost(profile=ANTICIPADO))
    assert r["status"] == S.REQUIERE_REVISION_TRIBUTARIA
    assert r["markup_pct"] is None and r["costs"]["reference"]["gross"] is None
    assert r["target_price"]["min"] is None and r["target_price"]["max"] is None
    assert r["costs"]["reference"]["gross_status"] == "requiere_revision_tributaria"
    assert r["reasons"] == ["impuesto_iva_anticipado"]


def test_unreliable_cost_is_qualified_even_if_net_verified():
    r = run(cost=cost(reliability=CostReliability.REQUIERE_REVISION), rule=rule(review_status=PolicyReviewStatus.VALIDADA), net_cost_verified=True)
    assert r["qualifiers"] == [S.COSTO_NO_VERIFICADO]
    assert r["validity"]["status"] == "provisional"


def test_validated_gross_cost_is_reported_as_unavailable():
    assert run()["costs"]["validated_gross"] == {"value": None, "status": "no_disponible"}


# ---------------------------------------------------------------- excepciones promocionales


def exc(**kw) -> PromotionException:
    base = dict(
        company_id=3,
        variant_id=10,
        price_list_id=12,
        kind=ExceptionKind.PROMOCION,
        starts_on=date(2026, 10, 1),
        ends_on=date(2026, 10, 31),
        applicable_gross_price=D("1300"),
        approval=ApprovalStatus.APROBADA,
    )
    base.update(kw)
    return PromotionException(**base)


def resolve(candidates, gross="1300", **kw) -> ExceptionResolution:
    args = dict(
        source_available=True,
        candidates=candidates,
        company_id=3,
        variant_id=10,
        price_list_id=12,
        bsale_gross_price=D(gross) if gross else None,
        today=TODAY,
    )
    args.update(kw)
    return resolve_exception(**args)


def test_unavailable_source_never_grants_exceptions():
    res = resolve([exc()], source_available=False)
    assert res.check is PromotionCheck.FUENTE_NO_INTEGRADA and res.verified_status is None


def test_exception_must_match_company_sku_list_dates_and_approval():
    assert resolve([exc(company_id=1)]).check is PromotionCheck.SIN_EXCEPCION
    assert resolve([exc(variant_id=11)]).check is PromotionCheck.SIN_EXCEPCION
    assert resolve([exc(price_list_id=13)]).check is PromotionCheck.SIN_EXCEPCION
    assert resolve([exc(ends_on=date(2026, 10, 12))]).check is PromotionCheck.SIN_EXCEPCION
    pending = resolve([exc(approval=ApprovalStatus.PENDIENTE)])
    assert pending.check is PromotionCheck.EXCEPCION_PENDIENTE_APROBACION
    assert pending.verified_status is None


def test_verified_exception_replaces_below_policy_but_keeps_cost_risks():
    res = resolve([exc()])
    assert res.check is PromotionCheck.EXCEPCION_VERIFICADA
    r = run(price=price("1300"), promotion=res, cost=cost(reliability=CostReliability.REQUIERE_REVISION))
    assert r["status"] == S.PROMOCION_VIGENTE
    assert S.COSTO_NO_VERIFICADO in r["qualifiers"]
    assert r["costs"]["reliability"] == "requiere_revision"
    remate = run(price=price("1000"), promotion=resolve([exc(kind=ExceptionKind.REMATE, applicable_gross_price=D("1000"))], gross="1000"))
    assert remate["status"] == S.REMATE_VIGENTE


def test_exception_with_price_mismatch_stays_pending_verification():
    res = resolve([exc()], gross="1350")
    assert res.check is PromotionCheck.EXCEPCION_NO_VERIFICADA
    r = run(price=price("1350"), promotion=res)
    assert r["status"] == S.BAJO_POLITICA
    assert S.PROMOCION_PENDIENTE_VERIFICACION in r["qualifiers"]


# ---------------------------------------------------------------- SQL


MARGIN_SQL = {n: v for n, v in vars(repo_mod).items() if n.endswith("_SQL") and isinstance(v, str)}


def test_margin_sql_reads_raw_and_only_margin_rules_outside_raw():
    for name, sql in MARGIN_SQL.items():
        assert_sql_is_read_only(sql)
        for schema, table in re.findall(r"\b(bsale\w*|app|analytics)\.(\w+)", sql):
            assert schema == "bsale_raw" or (schema, table) == ("bsale", "margin_rules"), (name, schema, table)
        assert "payload" not in sql
        assert "company_id = %s" in sql


# ---------------------------------------------------------------- servicio


class FakeMarginsRepo(MarginsV2Repository):
    def __init__(self, lists=None, prices=None, rules=None, stocks=None) -> None:
        self.calls: list[tuple[str, dict[str, Any]]] = []
        self._lists = lists if lists is not None else [
            {"price_list_id": 12, "name": "SUPERMERCADO", "state": 0, "missing_since": None},
            {"price_list_id": 13, "name": "RUTA Y WEB", "state": 0, "missing_since": None},
            {"price_list_id": 14, "name": "COMODITI ALCOHOLES", "state": 1, "missing_since": None},
            {"price_list_id": 16, "name": "MELINKA", "state": 0, "missing_since": None},
        ]
        self._prices = prices or []
        self._rules = rules or []
        self._stocks = stocks or []

    def price_lists(self, **kw):
        self.calls.append(("price_lists", kw))
        return self._lists

    def variant_prices(self, **kw):
        self.calls.append(("variant_prices", kw))
        ids = set(kw["price_list_ids"])
        return [p for p in self._prices if p["price_list_id"] in ids]

    def margin_rules(self, **kw):
        self.calls.append(("margin_rules", kw))
        return self._rules

    def stock_totals(self, **kw):
        self.calls.append(("stock_totals", kw))
        return self._stocks


def vp(list_id, variant_id, gross, missing=None):
    return {
        "price_list_id": list_id,
        "variant_id": variant_id,
        "variant_value": None,
        "variant_value_with_taxes": D(gross) if gross is not None else None,
        "missing_since": missing,
        "api_fetched_at": datetime(2026, 10, 13, tzinfo=timezone.utc),
    }


def mr(rid, list_id, type_id, min_m, max_m, active=True):
    return {
        "id": rid,
        "price_list_id": list_id,
        "price_list_name": f"L{list_id}",
        "product_type_id": type_id,
        "product_type_name": "ABARROTES",
        "min_margin": D(min_m),
        "max_margin": D(max_m),
        "active": active,
        "rule_key": f"3_{list_id}_{type_id}",
    }


STD_RULES = [mr(1, 12, 1, "20", "30"), mr(2, 13, 1, "10", "16"), mr(3, 14, 1, "5", "0"), mr(4, 16, 1, "0", "0")]


def make_service(facts, prices, rules=STD_RULES, lists=None, stocks=None) -> tuple[MarginsV2Service, FakeMarginsRepo, FakeRepo]:
    cost_repo = FakeRepo(facts=facts)
    costs = CostControlV2Service(cost_repo, SCOPE, today=TODAY, use_cache=False)
    if stocks is None:
        vids = {int(f["variant_id"]) for f in facts} | {int(p["variant_id"]) for p in prices}
        stocks = [{"variant_id": v, "stock_total": D("10"), "stock_offices": 1} for v in sorted(vids)]
    repo = FakeMarginsRepo(lists=lists, prices=prices, rules=rules, stocks=stocks)
    return MarginsV2Service(repo, costs, SCOPE, today=TODAY, use_cache=False), repo, cost_repo


def sku(variant_id, net="1000", **kw):
    row = fact(variant_id, average_cost=D("900"), **kw)
    if net is not None:
        row.update(purchase("lc_", variant_id, net, date(2026, 9, 1)))
        row.update(purchase("last_", variant_id, net, date(2026, 9, 1)))
    return row


def test_service_evaluates_only_vigent_authorized_lists_and_never_comoditi():
    svc, repo, _ = make_service([sku(10)], [vp(12, 10, "1500"), vp(13, 10, "1300"), vp(14, 10, "1100"), vp(16, 10, "1400")])
    out = svc.list_rows(cost_basis="ultimo_comparable", page_size=50)
    assert {r["price_list"]["id"] for r in out["items"]} == {12, 13, 16}
    price_call = next(kw for name, kw in repo.calls if name == "variant_prices")
    assert sorted(price_call["price_list_ids"]) == [12, 13, 16]
    assert all(kw["company_id"] == 3 for _, kw in repo.calls)
    with pytest.raises(MarginsValidationError):
        svc.list_rows(cost_basis="ultimo_comparable", price_list_id=14)


def test_service_status_per_list_and_rule_traceability():
    svc, _, _ = make_service([sku(10)], [vp(12, 10, "1500"), vp(13, 10, "1300"), vp(16, 10, "1400")])
    rows = {r["price_list"]["id"]: r for r in svc.list_rows(cost_basis="ultimo_comparable")["items"]}
    assert rows[12]["status"] == S.DENTRO_POLITICA
    assert rows[12]["rule"]["rule_key"] == "3_12_1" and rows[12]["rule"]["match"] == "tipo_producto"
    assert rows[13]["markup_pct"] == "9.24" and rows[13]["status"] == S.BAJO_POLITICA
    assert rows[16]["status"] == S.DENTRO_POLITICA
    assert rows[16]["rule"]["issues"] == [RuleIssue.MAXIMO_CERO]


def test_service_stock_total_by_variant_and_existence_filters():
    svc, _, _ = make_service(
        [sku(10), sku(11), sku(12), sku(13)],
        [vp(12, 10, "1500"), vp(13, 10, "1300"), vp(12, 11, "1500"), vp(12, 12, "1500"), vp(12, 13, "1500")],
        stocks=[
            {"variant_id": 10, "stock_total": D("12"), "stock_offices": 2},
            {"variant_id": 11, "stock_total": D("0"), "stock_offices": 6},
            {"variant_id": 12, "stock_total": D("-3"), "stock_offices": 6},
            # variante 13: sin filas sincronizadas
        ],
    )
    rows = svc.list_rows(cost_basis="ultimo_comparable", stock_filter="all", page_size=50)["items"]
    by = {(r["variant_id"], r["price_list"]["id"]): r["stock"] for r in rows}
    # El stock es de la variante: igual en sus dos listas, sin duplicar (12, no 24)
    assert by[(10, 12)]["total"] == "12" and by[(10, 13)]["total"] == "12"
    assert by[(10, 12)] == {"total": "12", "state": "con_stock", "offices": 2}
    assert by[(11, 12)]["state"] == "cero" and by[(11, 12)]["total"] == "0"
    assert by[(12, 12)]["state"] == "negativo" and by[(12, 12)]["total"] == "-3"
    assert by[(13, 12)] == {"total": None, "state": "sin_sincronizar", "offices": 0}

    con = svc.list_rows(cost_basis="ultimo_comparable", stock_filter="with", page_size=50)
    assert {r["variant_id"] for r in con["items"]} == {10} and con["total"] == 3  # sus 3 listas evaluadas
    assert con["stock_filter"] == "with"

    sin = svc.list_rows(cost_basis="ultimo_comparable", stock_filter="without", page_size=50)
    # Cero, negativo y sin fila (Fase 2.4: sin registro = SIN STOCK)
    assert {r["variant_id"] for r in sin["items"]} == {11, 12, 13}

    default = svc.list_rows(cost_basis="ultimo_comparable", page_size=50)
    assert default["stock_filter"] == "with"
    assert {r["variant_id"] for r in default["items"]} == {10}

    with pytest.raises(MarginsValidationError):
        svc.list_rows(cost_basis="ultimo_comparable", stock_filter="poco")


def test_margins_excel_includes_stock_total_column():
    from io import BytesIO

    from openpyxl import load_workbook

    from backend.services.analytics_v2 import excel_export

    svc, _, _ = make_service(
        [sku(10), sku(13)],
        [vp(12, 10, "1500"), vp(12, 13, "1500")],
        stocks=[{"variant_id": 10, "stock_total": D("12"), "stock_offices": 2}],
    )
    items = svc.list_rows(cost_basis="ultimo_comparable", stock_filter="all", page_size=50)["items"]
    blob = excel_export.margins_xlsx(
        meta={"company_name": "LA QUILLOTANA", "company_id": 3},
        sheets={"completo": items},
        filters_label="prueba",
    )
    wb = load_workbook(BytesIO(blob), read_only=True)
    ws = wb["Análisis completo"]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    assert "Stock Total" in headers
    col = headers.index("Stock Total")
    values = {r[5]: r[col] for r in rows[1:] if r[5] in (10, 13)}  # col 6 = ID de variante
    assert values[10] == 12
    assert values[13] is None  # sin sincronizar: vacío, nunca cero


def test_missing_price_and_missing_rule_are_reported():
    other_type = sku(11, product_type_id=99, product_type_name="SIN REGLA")
    svc, _, _ = make_service([sku(10), other_type], [vp(12, 10, "1500"), vp(12, 11, "1500")])
    rows = {(r["variant_id"], r["price_list"]["id"]): r for r in svc.list_rows(cost_basis="ultimo_comparable", page_size=200)["items"]}
    assert rows[(10, 13)]["status"] == S.SIN_PRECIO
    assert rows[(11, 12)]["status"] == S.SIN_REGLA
    review = svc.policy_review()
    assert {(m["price_list_id"], m["product_type_id"]) for m in review["missing_rules"]} == {(12, 99)}


def test_policy_review_keeps_retired_rules_as_history_without_alerts():
    svc, _, _ = make_service([sku(10)], [vp(12, 10, "1500")])
    review = svc.policy_review()
    comoditi = next(r for r in review["rules"] if r["price_list_id"] == 14)
    assert comoditi["in_commercial_scope"] is False and comoditi["alerts_enabled"] is False
    assert comoditi["price_list_scope"] == "retirada"
    assert review["policy"]["rules_total"] == 4
    assert review["policy"]["rules_in_scope"] == 3
    assert review["policy"]["rules_max_zero_in_scope"] == 1
    assert review["policy"]["rules_pending_review"] == 3


def test_summary_counts_by_list_and_reports_scope():
    svc, _, _ = make_service([sku(10), sku(11, net=None)], [vp(12, 10, "1500"), vp(12, 11, "1500")])
    out = svc.summary(cost_basis="ultimo_comparable")
    lst12 = next(b for b in out["by_price_list"] if b["price_list_id"] == 12)
    assert lst12["rows"] == 2 and lst12["by_status"][S.DENTRO_POLITICA] == 1 and lst12["by_status"][S.SIN_COSTO] == 1
    assert out["meta"]["policies_reviewed"] is False and out["meta"]["promotions_integrated"] is False
    assert out["scope"]["evaluated_price_list_ids"] == [12, 13, 16]
    # Compact rows del dashboard: tipo + barcode de la variante
    if out.get("top_alerts"):
        alert = out["top_alerts"][0]
        assert "barcode" in alert and "product_type" in alert


def test_reactivated_retired_list_is_reported_in_meta_and_not_evaluated():
    lists = [
        {"price_list_id": 12, "name": "A", "state": 0, "missing_since": None},
        {"price_list_id": 13, "name": "B", "state": 0, "missing_since": None},
        {"price_list_id": 14, "name": "COMODITI ALCOHOLES", "state": 0, "missing_since": None},
        {"price_list_id": 16, "name": "C", "state": 0, "missing_since": None},
    ]
    svc, _, _ = make_service([sku(10)], [vp(14, 10, "1100")], lists=lists)
    out = svc.summary(cost_basis="ultimo_comparable")
    assert out["meta"]["scope_discrepancies"] == 1
    assert out["scope"]["discrepancies"][0]["price_list_id"] == 14
    assert 14 not in {b["price_list_id"] for b in out["by_price_list"]}


def test_pagination_search_filters_and_validation_in_backend():
    facts = [sku(i, product_name=f"Producto {i}") for i in range(1, 121)]
    facts[4]["product_name"] = "Azúcar Iansa"
    prices = [vp(12, i, "1500") for i in range(1, 121)]
    svc, _, _ = make_service(facts, prices)
    page = svc.list_rows(cost_basis="ultimo_comparable", price_list_id=12, page=2, page_size=50)
    assert page["total"] == 120 and len(page["items"]) == 50 and page["total_pages"] == 3
    assert svc.list_rows(cost_basis="ultimo_comparable", q="azucar")["total"] == 3
    assert svc.list_rows(cost_basis="ultimo_comparable", status="sin_precio", price_list_id=13)["total"] == 120
    for bad in ({"status": "x"}, {"qualifier": "x"}, {"sort": "x"}, {"cost_basis": "x"}, {"variant_state": "x"}, {"reliability": "x"}):
        with pytest.raises(MarginsValidationError):
            svc.list_rows(**bad)


def test_sort_by_markup_puts_missing_values_last():
    svc, _, _ = make_service([sku(1), sku(2), sku(3, net=None)], [vp(12, 1, "1500"), vp(12, 2, "1300"), vp(12, 3, "1500")])
    items = svc.list_rows(cost_basis="ultimo_comparable", price_list_id=12, sort="markup_asc")["items"]
    assert [i["variant_id"] for i in items] == [2, 1, 3]


def test_product_detail_lists_every_evaluated_list_for_that_sku_only():
    svc, _, _ = make_service([sku(10), sku(11)], [vp(12, 10, "1500")])
    out = svc.product(10, cost_basis="ultimo_comparable")
    assert {r["price_list"]["id"] for r in out["rows"]} == {12, 13, 16}
    assert all(r["variant_id"] == 10 for r in out["rows"])


def test_performance_full_company_scale():
    n = 6000
    facts = [sku(i) for i in range(1, n + 1)]
    prices = [vp(lst, i, "1500") for lst in (12, 13, 16) for i in range(1, n + 1)]
    margins_service_mod.clear_cache()
    cost_service_mod.clear_cache()
    cost_repo = FakeRepo(facts=facts)
    costs = CostControlV2Service(cost_repo, SCOPE, today=TODAY, use_cache=True)
    repo = FakeMarginsRepo(prices=prices, rules=STD_RULES)
    svc = MarginsV2Service(repo, costs, SCOPE, today=TODAY, use_cache=True)
    try:
        start = time.perf_counter()
        first = svc.list_rows(cost_basis="ultimo_comparable", stock_filter="all", page_size=50)
        elapsed_first = time.perf_counter() - start
        start = time.perf_counter()
        svc.list_rows(
            cost_basis="ultimo_comparable",
            stock_filter="all",
            page=3,
            page_size=50,
            status="dentro_politica",
        )
        elapsed_cached = time.perf_counter() - start
    finally:
        margins_service_mod.clear_cache()
        cost_service_mod.clear_cache()
    assert first["total"] == n * 3 and len(first["items"]) == 50
    assert sum(1 for name, _ in repo.calls if name == "variant_prices") == 1
    assert elapsed_first < 8, elapsed_first
    assert elapsed_cached < 1, elapsed_cached
