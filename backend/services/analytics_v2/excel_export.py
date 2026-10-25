"""Exportadores Excel de Inteligencia Comercial (openpyxl, modo write-only).

Convenciones para que administración pueda cruzar datos contra Bsale:
- precios, costos y porcentajes se escriben como NÚMEROS de Excel (sin símbolos);
- SKU y códigos de barras como TEXTO (se preservan ceros y largos);
- fechas como fechas reales; celdas sin dato quedan vacías (nunca cero inventado);
- encabezados Graphite, primera fila congelada y filtros automáticos;
- cada archivo incluye la hoja «Información del archivo» con empresa, fecha,
  fuentes, filtros aplicados y estado de sincronización.
El modo write-only mantiene acotado el uso de memoria en exportaciones grandes.

Excel Premium (Auditoría Comercial 8.0): ver ``commercial_premium_xlsx``.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="17191E")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=12, color="17191E")
MONEY_FMT = '#,##0'
MONEY_DEC_FMT = '#,##0.00'
PCT_FMT = '0.00'
DATE_FMT = 'yyyy-mm-dd'
DATETIME_FMT = 'yyyy-mm-dd hh:mm'


def _num(value: Any) -> float | int | None:
    if value is None:
        return None
    try:
        d = Decimal(str(value))
    except Exception:
        return None
    return int(d) if d == d.to_integral_value() else float(d)


def _date(value: Any) -> date | datetime | None:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value
    text = str(value)
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        try:
            return date.fromisoformat(text[:10])
        except ValueError:
            return None


def _sheet(wb: Workbook, title: str, headers: Sequence[str], widths: Sequence[int]) -> Any:
    from openpyxl.cell import WriteOnlyCell

    ws = wb.create_sheet(title)
    # En write-only el congelado debe fijarse antes del primer append.
    ws.freeze_panes = "A2"
    header_cells: list[Any] = []
    for h in headers:
        cell = WriteOnlyCell(ws, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        header_cells.append(cell)
    ws.append(header_cells)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def _finish(ws: Any, rows_written: int, last_col: int) -> None:
    if rows_written > 0:
        ws.auto_filter.ref = f"A1:{get_column_letter(last_col)}{rows_written + 1}"


def _title_row(ws: Any, text: str) -> None:
    from openpyxl.cell import WriteOnlyCell

    cell = WriteOnlyCell(ws, value=text)
    cell.font = TITLE_FONT
    ws.append([cell])


def _info_sheet(wb: Workbook, title: str, entries: Sequence[tuple[str, Any]]) -> None:
    from openpyxl.cell import WriteOnlyCell

    ws = wb.create_sheet("Información del archivo")
    ws.freeze_panes = "A2"
    header_cells: list[Any] = []
    for h in ("Campo", "Detalle"):
        cell = WriteOnlyCell(ws, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        header_cells.append(cell)
    ws.append(header_cells)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    for key, value in entries:
        ws.append([key, "" if value is None else str(value)])


def _fmt_cells(ws: Any, row_idx: int, money_cols: set[int], pct_cols: set[int], date_cols: set[int]) -> None:
    # En write-only no se pueden releer celdas: los formatos se aplican al vuelo
    # en append mediante Cell con estilo; aquí solo queda como referencia.
    _ = (ws, row_idx, money_cols, pct_cols, date_cols)


def catalog_xlsx(
    *,
    rows: Iterable[dict[str, Any]],
    meta: dict[str, Any],
    list_ids: list[int],
    list_names: dict[int, str],
    scope_label: str,
    filters_label: str,
) -> tuple[bytes, int, set[str], set[str]]:
    """Catálogo Maestro. Retorna (bytes, filas, códigos_con_repetidos_avisos, barras_duplicadas)."""
    wb = Workbook(write_only=True)
    price_headers: list[str] = []
    for lid in list_ids:
        price_headers += [f"Precio bruto {list_names.get(lid, 'Lista ' + str(lid))} (#{lid})"]
    headers = [
        "Tipo de producto", "Producto", "Variante", "Código de barras", "SKU / código Bsale",
        "IVA %", "ILA %", "Impuestos", "Costo neto promedio Bsale", "Stock Total",
        *price_headers,
        "Estado", "company_id", "product_id", "variant_id",
        "Costo: última sincronización", "Precios: última sincronización",
        "Integridad", "Observaciones de integridad",
    ]
    widths = [22, 42, 30, 18, 18, 9, 9, 30, 18, 13] + [18] * len(list_ids) + [10, 11, 11, 11, 20, 20, 14, 30]
    ws = _sheet(wb, "Catálogo Maestro", headers, widths)

    from openpyxl.cell import WriteOnlyCell

    count = 0
    barcodes: dict[str, int] = {}
    codes: dict[str, int] = {}
    for r in rows:
        taxes_label = "; ".join(
            f"{t['name']} ({t['percentage']}%)" if t.get("percentage") is not None else t["name"]
            for t in r.get("taxes", [])
        ) or ("Sin configuración en Bsale" if r["tax_profile"]["status"] == "sin_impuestos" else "—")
        price_vals: list[Any] = []
        fetched_vals: list[str] = []
        for lid in list_ids:
            p = r["prices"].get(str(lid), {})
            price_vals.append(_num(p.get("gross")))
            if p.get("fetched_at"):
                fetched_vals.append(f"#{lid}: {str(p['fetched_at'])[:16].replace('T', ' ')}")
        values: list[Any] = [
            r.get("product_type"), r.get("product_name"), r.get("variant_name"),
            (str(r["barcode"]) if r.get("barcode") else None),
            (str(r["code"]) if r.get("code") else None),
            _num(r.get("iva_pct")), _num(r.get("ila_pct")), taxes_label,
            _num(r.get("average_cost_net")),
            _num(r.get("stock_total")),
            *price_vals,
            "Activa" if r["active"] else "Inactiva",
            r["company_id"], r["product_id"], r["variant_id"],
            _date(r.get("cost_fetched_at")),
            "; ".join(fetched_vals) or None,
            r["integrity_status"],
            ", ".join(r["integrity_issues"]) or None,
        ]
        cells: list[Any] = []
        money_idx = set(range(11, 11 + len(list_ids))) | {9}
        date_idx = {16 + len(list_ids)}
        text_idx = {4, 5}
        for i, v in enumerate(values, start=1):
            cell = WriteOnlyCell(ws, value=v)
            if i in money_idx:
                cell.number_format = MONEY_FMT
            elif i == 10:
                # Stock Total: cantidad (no dinero); vacío = sin sincronizar.
                cell.number_format = '#,##0.###'
            elif i in date_idx:
                cell.number_format = DATETIME_FMT
            elif i in (6, 7):
                cell.number_format = PCT_FMT
            if i in text_idx and v is not None:
                cell.number_format = "@"
            cells.append(cell)
        ws.append(cells)
        count += 1
        if r.get("barcode"):
            barcodes[str(r["barcode"])] = barcodes.get(str(r["barcode"]), 0) + 1
        if r.get("code"):
            codes[str(r["code"])] = codes.get(str(r["code"]), 0) + 1
    _finish(ws, count, len(headers))

    repeated_bar = sorted(k for k, n in barcodes.items() if n > 1)
    repeated_code = sorted(k for k, n in codes.items() if n > 1)
    freshness = meta.get("freshness", {})
    pl_freshness = {f["price_list_id"]: f for f in (meta.get("price_list_freshness") or [])}
    per_list_entries: list[tuple[str, Any]] = []
    for lid in list_ids:
        f = pl_freshness.get(lid)
        if f is None:
            per_list_entries.append((f"Lista #{lid}: sincronización de precios", "Sin registro de sincronización — PENDIENTE de confirmación"))
        else:
            detalle = f"estado {f.get('status') or '—'}"
            if f.get("last_success_at"):
                detalle += f" · último éxito {str(f['last_success_at'])[:16].replace('T', ' ')}"
            if f.get("last_error_at"):
                detalle += f" · último error {str(f['last_error_at'])[:16].replace('T', ' ')}"
            if f.get("rows_received") is not None:
                detalle += f" · filas recibidas {f['rows_received']}"
            if str(f.get("status") or "").upper() not in ("SUCCESS",):
                detalle += " — PENDIENTE de confirmación"
            per_list_entries.append((f"Lista #{lid}: sincronización de precios", detalle))
    _info_sheet(wb, "Información del archivo", [
        ("Archivo", "Catálogo Maestro — Inteligencia Comercial"),
        ("Alcance", scope_label),
        ("Filtros aplicados", filters_label),
        ("Empresa", f"{meta.get('company_name')} (company_id={meta.get('company_id')})"),
        ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Fuente", "BSALE RAW (solo lectura): productos, variantes, costos promedio, precios por lista e impuestos"),
        ("Listas incluidas", ", ".join(f"#{lid} {list_names.get(lid, '')}" for lid in list_ids)),
        ("Filas exportadas", count),
        ("Costo: última sincronización", freshness.get("variant_costs")),
        ("Stock: última sincronización", freshness.get("stocks")),
        ("Definición Stock Total", "Suma del stock físico actual (campo quantity de Bsale, bsale_raw.stocks) de la variante en todas las sucursales de la empresa. No es stock disponible ni reservado. Celda vacía = variante sin stock sincronizado (no equivale a cero); un 0 explícito sí es stock real en cero."),
        ("Precios: última sincronización", freshness.get("variant_prices")),
        *per_list_entries,
        ("Catálogo: última sincronización", freshness.get("products")),
        ("Impuestos: última sincronización", freshness.get("taxes")),
        ("Advertencia códigos de barras repetidos", (", ".join(repeated_bar[:30]) or "Sin códigos repetidos") + (" (la lista completa debe revisarse en la ERP)" if len(repeated_bar) > 30 else "")),
        ("Advertencia SKU repetidos", (", ".join(repeated_code[:30]) or "Sin SKU repetidos")),
        ("Nota", "Precios/costos son los sincronizados desde Bsale; celdas vacías significan dato ausente, no cero. Verificación contable pendiente."),
    ])
    if wb.sheetnames and wb.sheetnames[0] == "Sheet":
        del wb["Sheet"]
    out = BytesIO()
    wb.save(out)
    return out.getvalue(), count, set(repeated_code), set(repeated_bar)


def costs_xlsx(
    *,
    meta: dict[str, Any],
    summary: dict[str, Any],
    items: list[dict[str, Any]],
    priority: list[dict[str, Any]],
    reception_lines: list[dict[str, Any]],
    filters_label: str,
    reception_coverage: dict[str, Any] | None = None,
) -> bytes:
    from openpyxl.cell import WriteOnlyCell

    wb = Workbook(write_only=True)

    def money_cell(ws: Any, v: Any) -> Any:
        c = WriteOnlyCell(ws, value=_num(v))
        c.number_format = MONEY_FMT
        return c

    # Hoja 1: Resumen
    ws = wb.create_sheet("Resumen")
    _title_row(ws, "Resumen — Control de Costos")
    by_rel = summary.get("skus_active_by_reliability", {}) or {}
    rows_resumen = [
        ("Empresa", f"{meta.get('company_name')} (company_id={meta.get('company_id')})"),
        ("Fecha del análisis", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Fuente", "BSALE RAW (solo lectura)"),
        ("Total de SKU en alcance", summary.get("skus_active")),
        ("Costos consistentes", by_rel.get("consistente")),
        ("Requieren revisión", by_rel.get("requiere_revision")),
        ("Desactualizados", by_rel.get("desactualizado")),
        ("Sin costo", by_rel.get("sin_costo")),
        ("Solo promedio Bsale", by_rel.get("solo_promedio_bsale")),
        ("Filtros aplicados", filters_label),
        ("Estado de validación", "Costos e impuestos provisionales hasta validación contable; no sustituye la revisión administrativa."),
        (
            "Cobertura hoja «Recepciones relevantes»",
            (
                "COMPLETA para el universo filtrado"
                if not reception_coverage or reception_coverage.get("complete")
                else "PARCIAL — NO usar como historial completo: "
                + str(reception_coverage.get("reason") or "límite de exportación alcanzado")
            ),
        ),
    ]
    for k, v in rows_resumen:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 110

    product_headers = [
        "Tipo de producto", "Producto", "Variante", "SKU", "Código de barras",
        "Costo promedio neto", "Última compra comparable", "Compra comparable anterior",
        "Variación %", "Última recepción comparable", "Documento", "Sucursal",
        "Problema detectado", "Severidad", "Confiabilidad", "Observaciones para administración",
        "variant_id", "product_id",
    ]
    widths = [22, 40, 28, 16, 18, 16, 18, 18, 11, 16, 22, 18, 30, 11, 16, 42, 11, 11]

    def product_row(ws: Any, it: dict[str, Any]) -> list[Any]:
        lc = it.get("last_comparable") or {}
        pc = it.get("previous_comparable") or {}
        anomalies = it.get("anomalies") or []
        problema = "; ".join(a.get("label") or a.get("code", "") for a in anomalies) or None
        severidad = None
        for sev in ("alta", "media", "baja"):
            if any(a.get("severity") == sev for a in anomalies):
                severidad = sev
                break
        change = it.get("change") or {}
        vals: list[Any] = [
            (it.get("product_type") or {}).get("name"), it.get("product_name"), it.get("variant_name") or it.get("description"),
            (str(it["code"]) if it.get("code") else None), (str(it["bar_code"]) if it.get("bar_code") else None),
            money_cell(ws, (it.get("bsale_average_cost") or {}).get("value")),
            money_cell(ws, lc.get("cost")), money_cell(ws, pc.get("cost")),
            _num(change.get("pct")), _date(lc.get("admission_date")),
            (f"{lc.get('document') or ''} {lc.get('document_number') or ''}".strip() or None),
            lc.get("office_name") or (str(lc["office_id"]) if lc.get("office_id") is not None else None),
            problema, severidad, (it.get("reliability") or {}).get("status"),
            (it.get("reliability") or {}).get("reason"),
            it.get("variant_id"), it.get("product_id"),
        ]
        pct_cell = vals[8]
        if pct_cell is not None:
            c = WriteOnlyCell(ws, value=pct_cell)
            c.number_format = PCT_FMT
            vals[8] = c
        if isinstance(vals[9], (date, datetime)):
            c = WriteOnlyCell(ws, value=vals[9])
            c.number_format = DATE_FMT
            vals[9] = c
        for idx in (4, 5):
            if vals[idx - 1] is not None:
                c = WriteOnlyCell(ws, value=vals[idx - 1])
                c.number_format = "@"
                vals[idx - 1] = c
        return vals

    ws2 = _sheet(wb, "Productos prioritarios", product_headers, widths)
    for it in priority:
        ws2.append(product_row(ws2, it))
    _finish(ws2, len(priority), len(product_headers))

    ws3 = _sheet(wb, "Catálogo completo de costos", product_headers, widths)
    for it in items:
        ws3.append(product_row(ws3, it))
    _finish(ws3, len(items), len(product_headers))

    rec_headers = [
        "variant_id", "reception_id", "detail_id", "Fecha", "Sucursal (office_id)",
        "Movimiento", "Documento", "N° documento", "Cantidad", "Costo línea", "Comparable",
    ]
    ws4 = _sheet(wb, "Recepciones relevantes", rec_headers, [12, 13, 11, 13, 16, 18, 20, 14, 11, 14, 12])
    for l in reception_lines:
        ws4.append([
            l.get("variant_id"), l.get("reception_id"), l.get("detail_id"), _date(l.get("admission_date")),
            l.get("office_id"), l.get("movement_class"),
            l.get("document"), l.get("document_number"), _num(l.get("quantity")), _num(l.get("cost")),
            "Sí" if l.get("comparable") else "No",
        ])
    _finish(ws4, len(reception_lines), len(rec_headers))

    rev_headers = [
        "variant_id", "SKU", "Producto", "Estado de revisión", "Causa encontrada",
        "Acción recomendada", "Responsable", "Fecha de revisión", "Observaciones",
    ]
    ws5 = _sheet(wb, "Revisión administrativa", rev_headers, [12, 16, 40, 18, 30, 30, 18, 16, 44])
    for it in priority:
        ws5.append([
            it.get("variant_id"), (str(it["code"]) if it.get("code") else None), it.get("product_name"),
            None, None, None, None, None, None,
        ])
    _finish(ws5, len(priority), len(rev_headers))

    _info_sheet(wb, "Información del archivo", [
        ("Archivo", "Exportación de análisis — Control de Costos"),
        ("Empresa", f"{meta.get('company_name')} (company_id={meta.get('company_id')})"),
        ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Fuente", "BSALE RAW (solo lectura)"),
        ("Filtros aplicados", filters_label),
        ("SKU exportados", len(items)),
        ("SKU prioritarios", len(priority)),
        ("Líneas de recepción", len(reception_lines)),
        (
            "Cobertura de recepciones",
            (
                "Completa para el universo filtrado"
                if not reception_coverage or reception_coverage.get("complete")
                else "PARCIAL: " + str(reception_coverage.get("reason") or "límite alcanzado")
            ),
        ),
        ("Variantes del universo filtrado", (reception_coverage or {}).get("variants_total")),
        ("Variantes con recepciones incluidas", (reception_coverage or {}).get("variants_included")),
        ("Límite aplicado a recepciones", (reception_coverage or {}).get("limits")),
        (
            "Historial completo de recepciones",
            (reception_coverage or {}).get("segmented_hint")
            or "Si el universo supera el límite, descargue el historial por segmentos con el exportador independiente de recepciones (mismos filtros, parámetro segment).",
        ),
        ("Nota", "Las transferencias internas no se presentan como compras: se conserva la clasificación de movimiento de cada línea. Hoja de revisión para trabajo manual; aún no se importa a la ERP."),
    ])
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def margins_xlsx(
    *,
    meta: dict[str, Any],
    sheets: dict[str, list[dict[str, Any]]],
    filters_label: str,
) -> bytes:
    from openpyxl.cell import WriteOnlyCell

    wb = Workbook(write_only=True)

    ws = wb.create_sheet("Resumen")
    _title_row(ws, "Resumen — Análisis de Márgenes")

    for k, v in [
        ("Empresa", f"{meta.get('company_name')} (company_id={meta.get('company_id')})"),
        ("Fecha del análisis", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Fuente", "BSALE RAW (solo lectura) + política de márgenes vigente"),
        ("Filtros / modalidad", filters_label),
        ("Filas bajo costo", len(sheets.get("bajo_costo", []))),
        ("Filas bajo política", len(sheets.get("bajo_politica", []))),
        ("Filas análisis completo", len(sheets.get("completo", []))),
        ("Advertencia", "Recargo % se calcula sobre el costo; margen sobre venta % sobre el precio. El precio objetivo usa la fórmula vigente (costo bruto × (1 + mínimo)) solo cuando existen costo, impuestos y regla; en caso contrario queda vacío. Nada se modifica en Bsale."),
    ]:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 115

    headers = [
        "Tipo de producto", "Producto", "Variante", "Código de barras", "SKU", "ID de variante",
        "Lista de precios", "Precio bruto actual", "Costo neto promedio", "IVA %", "ILA %",
        "Costo bruto comercial", "Stock Total", "Recargo actual %", "Margen sobre venta %",
        "Recargo mínimo %", "Recargo máximo %", "Precio objetivo sugerido", "Diferencia vs actual",
        "Estado comercial", "Confiabilidad del costo", "Advertencia tributaria", "Estado de la política",
    ]
    widths = [22, 40, 28, 18, 16, 12, 20, 17, 16, 9, 9, 17, 13, 13, 14, 13, 13, 18, 15, 20, 18, 24, 20]

    def row_cells(w: Any, r: dict[str, Any]) -> list[Any]:
        costs = r.get("costs") or {}
        ref = costs.get("reference") or {}
        rule = r.get("rule") or {}
        price = (r.get("price") or {}).get("gross")
        # Precio objetivo oficial del motor (solo existe con costo, impuestos y
        # regla usable; en cualquier otro caso queda vacío, no se inventa).
        target = (r.get("target_price") or {}).get("min")
        diff = None
        if target is not None and price is not None:
            diff = Decimal(str(target)) - Decimal(str(price))
        vals: list[Any] = [
            (r.get("product_type") or {}).get("name"), r.get("product_name"), r.get("variant_name"),
            (str(r["barcode"]) if r.get("barcode") else None), (str(r["code"]) if r.get("code") else None),
            r.get("variant_id"), (r.get("price_list") or {}).get("label") or (r.get("price_list") or {}).get("name"),
            _num(price), _num(ref.get("net")), _num(ref.get("iva_rate_pct") if "iva_rate_pct" in ref else None),
            _num(ref.get("ila_rate_pct") if "ila_rate_pct" in ref else None), _num(ref.get("gross")),
            # Stock Total: cantidad física acumulada (vacío = sin sincronizar).
            _num((r.get("stock") or {}).get("total")),
            _num(r.get("markup_pct")), _num(r.get("margin_on_sale_pct")),
            _num(rule.get("min_markup_pct")), _num(rule.get("max_markup_pct")),
            _num(target), _num(diff), r.get("status"), costs.get("reliability"),
            ref.get("tax_status"), (r.get("validity") or {}).get("status"),
        ]
        out: list[Any] = []
        for i, v in enumerate(vals, start=1):
            c = WriteOnlyCell(w, value=v)
            if i in (8, 9, 12, 18, 19):
                c.number_format = MONEY_FMT
            elif i == 13:
                c.number_format = '#,##0.###'
            elif i in (10, 11, 14, 15, 16, 17):
                c.number_format = PCT_FMT
            elif i in (4, 5):
                c.number_format = "@"
            out.append(c)
        return out

    sheet_titles = {
        "bajo_costo": "Productos bajo costo",
        "bajo_politica": "Productos bajo política",
        "completo": "Análisis completo",
    }
    for key, title in sheet_titles.items():
        if key not in sheets:
            continue
        rows = sheets[key]
        w = _sheet(wb, title, headers, widths)
        for r in rows:
            w.append(row_cells(w, r))
        _finish(w, len(rows), len(headers))

    rev_headers = ["ID de variante", "SKU", "Producto", "Lista de precios", "Conclusión", "Observaciones", "Propuesta", "Responsable", "Fecha"]
    w = _sheet(wb, "Revisión administrativa", rev_headers, [13, 16, 40, 20, 30, 36, 30, 18, 13])
    seen: set[tuple[Any, Any]] = set()
    revision_source = [r for rows in sheets.values() for r in rows]
    for r in revision_source:
        key = (r.get("variant_id"), (r.get("price_list") or {}).get("id"))
        if key in seen:
            continue
        seen.add(key)
        w.append([r.get("variant_id"), (str(r["code"]) if r.get("code") else None), r.get("product_name"),
                  (r.get("price_list") or {}).get("label") or (r.get("price_list") or {}).get("name"),
                  None, None, None, None, None])
    _finish(w, len(seen), len(rev_headers))

    w = wb.create_sheet("Metodología y validación")
    _title_row(w, "Metodología y estado de validación")

    for line in [
        "Recargo % = (Precio bruto / Costo bruto − 1) × 100. Margen sobre venta % = (Precio − Costo) / Precio × 100. No son equivalentes.",
        "Precio objetivo mínimo = Costo bruto comercial × (1 + recargo mínimo de política / 100), redondeado al peso superior (fórmula vigente del motor comercial).",
        "El costo bruto comercial y los márgenes son PROVISIONALES hasta la validación contable de costos netos e impuestos (IVA/ILA, IVA anticipado).",
        "Sin costo, sin impuestos identificables o sin regla de política, el precio objetivo queda vacío: no se inventan recomendaciones.",
        "Fuentes: BSALE RAW (precios, costos promedio, impuestos), política de márgenes vigente y excepciones de promociones verificadas.",
        "Este archivo es de análisis: no modifica precios en Bsale ni aprueba cambios automáticamente.",
    ]:
        w.append([line])
    w.column_dimensions["A"].width = 150

    _info_sheet(wb, "Información del archivo", [
        ("Archivo", "Exportación de análisis — Márgenes"),
        ("Empresa", f"{meta.get('company_name')} (company_id={meta.get('company_id')})"),
        ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Fuente", "BSALE RAW (solo lectura)"),
        ("Filtros / modalidad", filters_label),
    ])
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def receptions_xlsx(
    *,
    meta: dict[str, Any],
    reception_lines: list[dict[str, Any]],
    filters_label: str,
    segment: int,
    segment_size: int,
    variants_total: int,
    variants_included: int,
    segments_total: int,
    lines_truncated: bool,
) -> bytes:
    """Historial de recepciones por segmento de variantes (exportación independiente).

    Nunca se presenta como completo si el segmento alcanzó el tope de líneas:
    queda declarado en la hoja de información junto al segmento, los filtros y
    los límites aplicados.
    """
    wb = Workbook(write_only=True)
    headers = [
        "variant_id", "SKU", "Producto", "Variante", "reception_id", "detail_id",
        "Fecha", "Sucursal (office_id)", "Movimiento", "Documento", "N° documento",
        "Cantidad", "Costo línea", "Comparable",
    ]
    ws = _sheet(wb, "Recepciones", headers, [12, 16, 40, 28, 13, 11, 13, 16, 18, 20, 14, 11, 14, 12])
    for l in reception_lines:
        ws.append([
            l.get("variant_id"),
            (str(l["code"]) if l.get("code") else None),
            l.get("product_name"), l.get("variant_name"),
            l.get("reception_id"), l.get("detail_id"), _date(l.get("admission_date")),
            l.get("office_id"), l.get("movement_class"),
            l.get("document"), l.get("document_number"), _num(l.get("quantity")), _num(l.get("cost")),
            "Sí" if l.get("comparable") else "No",
        ])
    _finish(ws, len(reception_lines), len(headers))

    _info_sheet(wb, "Información del archivo", [
        ("Archivo", "Historial de recepciones — exportación segmentada"),
        ("Empresa", f"{meta.get('company_name')} (company_id={meta.get('company_id')})"),
        ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Fuente", "BSALE RAW (solo lectura): stock_receptions + stock_reception_details (líneas y costos originales, sin recalcular)"),
        ("Filtros aplicados", filters_label),
        ("Segmento", f"{segment} de {segments_total}"),
        ("Tamaño de segmento (variantes)", segment_size),
        ("Variantes del universo filtrado", variants_total),
        ("Variantes incluidas en este segmento", variants_included),
        ("Líneas incluidas", len(reception_lines)),
        (
            "Cobertura",
            (
                "PARCIAL dentro del segmento: se alcanzó el tope de líneas; reduzca el tamaño de segmento para completar este tramo."
                if lines_truncated
                else "Completa para las variantes de este segmento. El historial total requiere descargar todos los segmentos."
            ),
        ),
        ("Nota", "Las transferencias internas conservan su clasificación y no se presentan como compras. No se fusionan ni modifican recepciones."),
    ])
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def commercial_premium_xlsx(payload: dict[str, Any]) -> bytes:
    """Excel Premium Auditoría Comercial 8.0 (6 hojas). Reexporta el generador dedicado."""
    from backend.services.analytics_v2.commercial_premium_excel import (
        commercial_premium_xlsx as _build,
    )

    return _build(payload)


def commercial_premium_review_xlsx(payload: dict[str, Any]) -> bytes:
    """Excel operativo de revisión (2 hojas). Mismo payload; reexporta el generador dedicado."""
    from backend.services.analytics_v2.commercial_premium_excel import (
        commercial_premium_review_xlsx as _build,
    )

    return _build(payload)
