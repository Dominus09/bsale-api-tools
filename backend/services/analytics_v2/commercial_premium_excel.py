"""Generador XLSX — Excel Premium Auditoría Comercial 8.0 (6 hojas).

Extiende el estilo de ``excel_export`` (Graphite, filtros, freeze, texto SKU/barras).
No escribe PostgreSQL; la hoja de revisión es solo plantilla editable local.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from backend.services.analytics_v2.commercial_premium_service import REVIEW_SHEET_NAMES, SHEET_NAMES
from backend.services.analytics_v2.display_labels import office_visible_name, price_column_header
from backend.services.analytics_v2.product_identity import (
    format_barcode,
    format_product_plus_variant,
    format_product_type,
)
from backend.services.analytics_v2.stock_scope import STOCK_CRITERION, STOCK_FILTER_LABELS

# Paleta alineada a Excel PRO aprobado por administración.
BANNER_FILL = PatternFill("solid", fgColor="131C2B")
BANNER_SUB_FILL = PatternFill("solid", fgColor="202D3E")
HEADER_FILL = PatternFill("solid", fgColor="202D3E")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=9)
BANNER_FONT = Font(color="FFFFFF", bold=True, size=18)
BANNER_SUB_FONT = Font(color="FFFFFF", bold=False, size=10)
INK_FONT = Font(color="1B2637", bold=False, size=10)
INK_BOLD = Font(color="1B2637", bold=True, size=10)
MUTED_FONT = Font(color="536276", bold=False, size=9)
FILL_ZEBRA = PatternFill("solid", fgColor="F8FAFD")
FILL_CHIP = PatternFill("solid", fgColor="F1F5FA")
FILL_STOCK_OK = PatternFill("solid", fgColor="ECF8F5")
FONT_STOCK_OK = Font(color="19735D", bold=True, size=11)
DATE_FMT = "dd-mm-yyyy"
DATETIME_FMT = "yyyy-mm-dd hh:mm"
PCT_FMT = "0.00"


def _num(value: Any) -> float | int | None:
    if value is None:
        return None
    try:
        d = Decimal(str(value))
    except Exception:
        return None
    return int(d) if d == d.to_integral_value() else float(d)


def _date(value: Any) -> date | datetime | None:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value
    text = str(value)
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        try:
            return date.fromisoformat(text[:10])
        except ValueError:
            return None

# Pesos chilenos (formato PRO) y variaciones % con signo.
CLP_FMT = '"$"\\ #,##0;[Red]\\("$"\\ #,##0\\);"—"'
CLP_DEC_FMT = '"$"\\ #,##0.00;[Red]\\("$"\\ #,##0.00\\);"—"'
QTY_FMT = '#,##0;[Red]\\(#,##0\\);"—"'
PCT_RATIO_FMT = '\\+0.0%;[Red]\\-0.0%;"—"'
THIN = Border(
    left=Side(style="thin", color="D0D4DA"),
    right=Side(style="thin", color="D0D4DA"),
    top=Side(style="thin", color="D0D4DA"),
    bottom=Side(style="thin", color="D0D4DA"),
)
FILL_P1 = PatternFill("solid", fgColor="FFF0F2")
FILL_P2 = PatternFill("solid", fgColor="FFF7E8")
FILL_P3 = PatternFill("solid", fgColor="EEF4FE")
FILL_PENDING = PatternFill("solid", fgColor="EEF4FE")
FILL_WARN = PatternFill("solid", fgColor="FFF7E8")
FILL_SECTION = PatternFill("solid", fgColor="202D3E")
FONT_SECTION = Font(color="FFFFFF", bold=True, size=11)
FONT_MUTED = Font(color="536276", size=9, italic=True)
FONT_P1 = Font(color="E3485C", bold=True, size=9)
FONT_P2 = Font(color="E6A540", bold=True, size=9)
FONT_P3 = Font(color="3158A6", bold=True, size=9)
FONT_EDITABLE = Font(color="3158A6", bold=True, size=10)
ADMIN_STATES = (
    "Pendiente",
    "En revisión",
    "Revisado sin error",
    "Corrección necesaria",
    "Corrección aplicada",
)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    return str(value)


def _priority_fill(band: str | None) -> PatternFill | None:
    if not band:
        return None
    b = band.upper()
    if b.startswith("P1"):
        return FILL_P1
    if b.startswith("P2"):
        return FILL_P2
    if b.startswith("P3"):
        return FILL_P3
    return None


def _priority_label(band: str | None) -> str | None:
    if not band:
        return None
    b = str(band).upper()
    if b.startswith("P1"):
        return "P1 · CRÍTICO"
    if b.startswith("P2"):
        return "P2 · ALTO"
    if b.startswith("P3"):
        return "P3 · MEDIO"
    return str(band)


def _priority_font(band: str | None) -> Font | None:
    if not band:
        return None
    b = str(band).upper()
    if b.startswith("P1"):
        return FONT_P1
    if b.startswith("P2"):
        return FONT_P2
    if b.startswith("P3"):
        return FONT_P3
    return None


def _merge_row(ws: Any, row: int, cols: int, value: Any, *, fill: PatternFill, font: Font) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, value)
    cell.fill = fill
    cell.font = font
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    for c in range(1, cols + 1):
        ws.cell(row, c).fill = fill


def _banner_pair(ws: Any, *, title: str, subtitle: str, cols: int = 11) -> None:
    _merge_row(ws, 1, cols, title, fill=BANNER_FILL, font=BANNER_FONT)
    _merge_row(ws, 2, cols, subtitle, fill=BANNER_SUB_FILL, font=BANNER_SUB_FONT)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18


def _set_widths(ws: Any, widths: Sequence[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _header_row(ws: Any, headers: Sequence[str], row: int = 1) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = THIN
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _style_data_cell(
    cell: Any,
    *,
    money: bool = False,
    money_dec: bool = False,
    pct_ratio: bool = False,
    pct_points: bool = False,
    as_date: bool = False,
    as_datetime: bool = False,
    as_text: bool = False,
    fill: PatternFill | None = None,
    zebra: bool = False,
) -> None:
    cell.border = THIN
    cell.alignment = Alignment(vertical="center")
    if fill is not None:
        cell.fill = fill
    elif zebra:
        cell.fill = FILL_ZEBRA
    if as_text:
        cell.number_format = "@"
        if cell.value is not None and not isinstance(cell.value, str):
            cell.value = str(cell.value)
        # Evitar que Excel interprete strings que empiezan con = como fórmula.
        if isinstance(cell.value, str) and cell.value[:1] in ("=", "+", "-", "@"):
            cell.value = "'" + cell.value
    elif money_dec:
        cell.number_format = CLP_DEC_FMT
    elif money:
        cell.number_format = CLP_FMT
    elif pct_ratio:
        cell.number_format = PCT_RATIO_FMT
    elif pct_points:
        cell.number_format = PCT_FMT
    elif as_datetime:
        cell.number_format = DATETIME_FMT
    elif as_date:
        cell.number_format = DATE_FMT


def _office_label(payload: dict[str, Any], office_id: Any) -> str | None:
    if office_id is None:
        return None
    offices: dict[Any, str] = payload.get("offices") or {}
    company_id = payload.get("company_id")
    try:
        key = int(office_id)
    except (TypeError, ValueError):
        return office_visible_name(office_id=None)
    if key in offices:
        return offices[key]
    try:
        cid = int(company_id) if company_id is not None else None
    except (TypeError, ValueError):
        cid = None
    return office_visible_name(office_id=key, company_id=cid)


def _price_or_missing(value: Any) -> Any:
    """Marca de presentación para precio ausente. Nunca se usa en fórmulas numéricas."""
    n = _num(value)
    return n if n is not None else "Sin precio"


def _table_header(ws: Any, headers: Sequence[str], row: int) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = THIN


def _write_summary(ws: Any, payload: dict[str, Any]) -> None:
    ws.sheet_view.showGridLines = False
    _set_widths(ws, [47, 56, 55])
    _banner_pair(
        ws,
        title="Excel Premium — Auditoría Comercial 8.0",
        subtitle="Panel administrativo de control de costos · alertas ≠ errores confirmados",
        cols=3,
    )
    ws.cell(1, 1).font = Font(color="FFFFFF", bold=True, size=19)

    census = payload.get("census") or {}
    costs = payload.get("cost_coverage") or {}
    buckets = payload.get("divergence_buckets") or {}
    zsum = payload.get("zero_cost_summary") or {}
    identity = payload.get("divergence_buckets_identity") or {}
    stock_filt = payload.get("stock_filter") or "with"
    stock_label = payload.get("stock_filter_label") or STOCK_FILTER_LABELS.get(stock_filt, stock_filt)

    rows: list[tuple[str, Any, str]] = [
        ("FILTRO DE STOCK", stock_label, payload.get("stock_criterion") or STOCK_CRITERION),
        ("Empresa", f"{payload.get('company_name')} (company_id={payload.get('company_id')})", ""),
        ("Corte UTC", payload.get("cut_at_utc"), "Fotografía REPEATABLE READ"),
        ("Indicadores globales", "census / cobertura / huérfanas", "Sin filtro de stock operativo"),
        ("Variantes totales", census.get("variantes"), "GLOBAL"),
        ("Activas / Inactivas", f"{census.get('activas')} / {census.get('inactivas')}", "GLOBAL"),
        ("Cobertura costos (filas)", costs.get("filas_variant_costs"), f"GLOBAL · de {costs.get('variantes')} variantes"),
        ("Costos promedio > 0", costs.get("costs_average_pos"), "GLOBAL · variant_costs.average_cost"),
        ("Costos promedio = 0", costs.get("costs_average_zero"), "GLOBAL · no es ausencia de fila"),
        ("Activas stock>0", census.get("activas_stock_pos"), "GLOBAL"),
        ("Huérfanas stock>0 sin variante", census.get("stock_pos_sin_variante"), "GLOBAL · integridad"),
        ("Universo operativo (filtro stock)", buckets.get("universo"), f"Promedio vs recepciones · {stock_label}"),
        ("Divergencias ≤15 %", buckets.get("hasta_15"), "Universo filtrado"),
        ("Divergencias 15–50 % (moderadas)", buckets.get("entre_15_y_50"), "Universo filtrado"),
        ("Divergencias >50 % (importantes)", buckets.get("sobre_50"), "Universo filtrado"),
        ("Indicador extremas (≥100 % o ≤−60 %)", buckets.get("extrema_indicador"), "Subconjunto de >50 %; no error confirmado"),
        ("Requieren verificación de evidencia", buckets.get("requiere_verificacion"), "Universo filtrado"),
        ("Sin recepción comparable", buckets.get("sin_recepcion_comparable"), "Universo filtrado"),
        ("Anomalías exportadas (Δ>15 %)", payload.get("anomalies_count"), payload.get("anomalies_source")),
        ("Activos costo cero (filtrados)", zsum.get("total"), zsum.get("scope") or ""),
        ("  · stock cero explícito", zsum.get("stock_cero_explicito"), "Fila stock con quantity=0"),
        ("  · sin fila de stock", zsum.get("sin_fila_stock"), "Dato ausente ≠ cero"),
        ("Cierre buckets principales", "Sí" if identity.get("closes") else "NO", str(identity)),
    ]

    r = 4
    _table_header(ws, ("INDICADOR", "DATO DEL CORTE", "ALCANCE / OBSERVACIÓN"), r)
    r = 5
    for idx, (label, value, note) in enumerate(rows):
        zebra = idx % 2 == 1
        for col, val in enumerate((label, value if value is not None else "—", note), start=1):
            cell = ws.cell(r, col, val)
            _style_data_cell(cell, zebra=zebra)
            if col == 1:
                cell.font = INK_FONT
            elif col == 3:
                cell.font = MUTED_FONT
        r += 1

    r += 1
    sec = ws.cell(r, 1, "Metodología y advertencias comerciales")
    sec.fill = FILL_SECTION
    sec.font = FONT_SECTION
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    r += 1
    method_lines = [
        "Costo analizado = average_cost Bsale (no se recalcula desde la última recepción).",
        "Δ% = (última comparable − promedio) / promedio × 100. Comparable = FACTURA/GUÍA completa con costo>0 y cantidad>0.",
        "Exposición potencial = |última comparable − promedio| × stock. Solo orientación de revisión; no es pérdida ni valorización contable.",
        "Si la unidad/presentación no está comprobada, la exposición no debe usarse como valorización definitiva.",
        "Prioridad de revisión (P1–P4) es independiente de la confianza de error.",
        "Una posible corrección administrativa no reduce la prioridad de revisar el caso.",
        "Alertas ≠ errores confirmados. Editar la hoja 6 no actualiza la ERP ni PostgreSQL.",
        "«Sin precio» es solo presentación; internamente el precio ausente permanece NULL y no entra en cálculos.",
    ]
    for line in method_lines + list(payload.get("warnings") or []):
        ws.cell(r, 1, "• " + line).font = Font(size=9)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        r += 1

    r += 1
    sec2 = ws.cell(r, 1, "Frescura de sincronización (solo alertas activas)")
    sec2.fill = FILL_SECTION
    sec2.font = FONT_SECTION
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    r += 1
    _table_header(ws, ("Recurso", "Estado / último SUCCESS", "Nota"), r)
    r += 1
    for s in payload.get("sync_state") or []:
        ws.cell(r, 1, f"{s.get('resource')} / {s.get('scope')}")
        success = s.get("last_success_at")
        ws.cell(r, 2, f"{s.get('status')} · {success}")
        note_cell = ws.cell(r, 3, s.get("freshness_note"))
        if s.get("active_freshness_alert"):
            note_cell.fill = FILL_P2
        r += 1
    ws.freeze_panes = "A5"


def _write_anomalies(ws: Any, payload: dict[str, Any]) -> None:
    """Críticos y anomalías — layout PRO + precios por nombre comercial."""
    price_lists: dict[int, str] = payload.get("price_lists") or {}
    list_ids = sorted(price_lists.keys())
    headers = [
        "TIPO",
        "PRODUCTO + VARIANTE",
        "BARCODE",
        "PRIORIDAD",
        "STOCK",
        "COSTO PROM.",
        "ÚLT. COMPRA",
        "Δ vs PROM.",
        "EXPOSICIÓN¹",
        "ACCIÓN",
        "N° FACTURA",
        "FECHA",
        "OFICINA",
    ]
    for lid in list_ids:
        headers.append(price_column_header(str(price_lists.get(lid) or f"lista {lid}")))
    headers.extend(
        [
            "VARIANT_ID",
            "SKU",
            "SCORE",
            "TIPO ALERTA",
            "EVIDENCIA",
            "CONFIANZA ERROR",
            "DOCUMENTO",
            "NOTA METODOLÓGICA",
        ]
    )
    n_base = 13
    price_col0 = n_base + 1
    price_col_last = n_base + len(list_ids)
    tech_sku_col = price_col_last + 2

    rows = payload.get("anomalies") or []
    stock_label = payload.get("stock_filter_label") or "CON STOCK"
    cols = len(headers)
    _set_widths(
        ws,
        [17, 54, 19, 16, 11, 16, 14, 13, 18, 34, 15, 14, 18]
        + [15] * len(list_ids)
        + [14, 16, 11, 25, 14, 16, 14, 40],
    )
    _banner_pair(
        ws,
        title="ANOMALÍAS DE COSTO  /  AUDITORÍA COMPLETA",
        subtitle=(
            f"{len(rows)} diferencias superiores al 15 %  ·  {stock_label}  ·  "
            "Alerta ≠ error confirmado"
        ),
        cols=cols,
    )
    _merge_row(
        ws,
        3,
        cols,
        "PRIORIDAD  →  PRODUCTO + VARIANTE  →  COSTO / Δ %  →  FACTURA.  "
        "Columnas técnicas al final. «Sin precio» es solo presentación (no es cero).",
        fill=FILL_CHIP,
        font=MUTED_FONT,
    )
    header_row = 5
    _table_header(ws, headers, header_row)
    data_start = 6
    for offset, item in enumerate(rows):
        i = data_start + offset
        prio = item.get("admin_priority") or {}
        band = prio.get("review_priority") or prio.get("priority_band")
        fill = _priority_fill(band)
        zebra = offset % 2 == 1
        diff = _num(item.get("diff_pct"))
        diff_ratio = (float(diff) / 100.0) if diff is not None else None
        doc = item.get("last_comparable_document") or item.get("lc_document")
        doc_n = item.get("last_comparable_document_number") or item.get("lc_document_number")
        office_id = item.get("last_comparable_office_id") or item.get("lc_office_id")
        values: list[Any] = [
            format_product_type(item.get("product_type_name")),
            format_product_plus_variant(item.get("product_name"), item.get("variant_name")),
            format_barcode(item.get("barcode") or item.get("sku")),
            _priority_label(band),
            _num(item.get("stock_total")),
            _num(item.get("average_cost")),
            _num(item.get("last_comparable_cost")),
            diff_ratio,
            None,
            prio.get("suggested_action"),
            _text(doc_n),
            _date(item.get("last_comparable_date") or item.get("lc_date")),
            _office_label(payload, office_id),
        ]
        for lid in list_ids:
            # Presentación: «Sin precio». La fórmula de exposición usa E/F/G, nunca estas columnas.
            values.append(_price_or_missing(item.get(f"price_{lid}")))
        values.extend(
            [
                item.get("variant_id"),
                _text(item.get("sku")),
                prio.get("priority_score"),
                "divergencia_gt_15_alerta" + ("_extrema" if item.get("extrema_indicador") else ""),
                prio.get("evidence_quality"),
                prio.get("error_confidence"),
                doc,
                prio.get("monetary_exposure_disclaimer"),
            ]
        )
        for col, val in enumerate(values, start=1):
            cell = ws.cell(i, col, val)
            is_price = price_col0 <= col <= price_col_last
            missing_price = is_price and val == "Sin precio"
            _style_data_cell(
                cell,
                money_dec=col in (6, 7) or (is_price and not missing_price),
                pct_ratio=col == 8,
                as_date=col == 12,
                as_text=col in (3, 11) or missing_price or col == tech_sku_col,
                fill=fill if col in (4, 8) else None,
                zebra=zebra and fill is None,
            )
            if col == 4:
                pf = _priority_font(band)
                if pf:
                    cell.font = pf
            if col == 5 and cell.value is not None:
                cell.number_format = QTY_FMT
        formula = f'=IF(OR(E{i}="",F{i}="",G{i}=""),"",ABS(G{i}-F{i})*E{i})'
        exp_cell = ws.cell(i, 9, formula)
        _style_data_cell(exp_cell, money_dec=True, fill=FILL_PENDING, zebra=zebra)

    last = data_start + max(len(rows), 1) - 1 if rows else header_row
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(cols)}{last}"
    ws.freeze_panes = f"D{data_start}"


def _write_zero_cost(ws: Any, payload: dict[str, Any]) -> None:
    price_lists: dict[int, str] = payload.get("price_lists") or {}
    list_ids = sorted(price_lists.keys())
    headers = [
        "Producto",
        "Variante",
        "variant_id",
        "SKU",
        "Código de barras",
        "Costo promedio",
        "Estado stock",
        "Stock total (vacío=sin fila)",
        "Stock cero explícito",
        "Sin fila stock",
        "Con recepción comparable",
        "Sin recepción comparable",
        "Estado comercial",
        "Acción recomendada",
        "Documento comparable",
        "Fecha comparable",
        "Oficina",
    ]
    for lid in list_ids:
        headers.append(price_column_header(str(price_lists.get(lid) or f"lista {lid}")))

    rows = payload.get("zero_cost_actives") or []
    cols = len(headers)
    stock_label = payload.get("stock_filter_label") or "CON STOCK"
    _set_widths(ws, [25, 28, 11, 16, 16, 12, 18, 14, 12, 12, 12, 12, 22, 32, 18, 13, 18] + [16] * len(list_ids))
    _banner_pair(
        ws,
        title="ACTIVOS CON COSTO CERO",
        subtitle=(
            f"Alcance operativo {stock_label}. "
            + (
                f"{len(rows)} registros en este corte."
                if rows
                else "No se identificaron registros en este corte."
            )
        ),
        cols=cols,
    )
    note = (
        "Costo promedio = 0 explícito. Stock vacío = sin fila (dato ausente ≠ cero). "
        "«Sin precio» es presentación; no entra en cálculos."
    )
    _merge_row(
        ws,
        3,
        cols,
        note,
        fill=FILL_STOCK_OK if not rows else FILL_WARN,
        font=FONT_STOCK_OK if not rows else INK_FONT,
    )
    header_row = 5
    _table_header(ws, headers, header_row)
    data_start = 6
    price_col0 = 18
    for offset, z in enumerate(rows):
        i = data_start + offset
        sin_fila = bool(z.get("sin_fila_stock") or z.get("stock_status") == "sin_fila_stock")
        stock_val = None if sin_fila else _num(
            z.get("stock_total_raw") if z.get("stock_total_raw") is not None else z.get("stock_total")
        )
        values: list[Any] = [
            z.get("product_name"),
            z.get("variant_name"),
            z.get("variant_id"),
            _text(z.get("sku")),
            _text(z.get("barcode")),
            _num(z.get("average_cost")),
            z.get("stock_status"),
            stock_val,
            "Sí" if z.get("stock_status") == "stock_cero_explicito" else "No",
            "Sí" if sin_fila else "No",
            "Sí" if z.get("has_comparable_reception") else "No",
            "No" if z.get("has_comparable_reception") else "Sí",
            z.get("commercial_status"),
            z.get("recommended_action"),
            z.get("last_comparable_document"),
            _date(z.get("last_comparable_date")),
            _office_label(payload, z.get("last_comparable_office_id")),
        ]
        for lid in list_ids:
            values.append(_price_or_missing(z.get(f"price_{lid}")))

        pending = sin_fila or not z.get("has_comparable_reception")
        zebra = not pending and (offset % 2 == 1)
        for col, val in enumerate(values, start=1):
            cell = ws.cell(i, col, val)
            is_price = col >= price_col0
            missing_price = is_price and val == "Sin precio"
            _style_data_cell(
                cell,
                money=(col == 6 or (is_price and not missing_price)),
                as_date=col == 16,
                as_text=col in (4, 5) or missing_price,
                fill=FILL_PENDING if pending and col in (7, 8) else None,
                zebra=zebra,
            )
    last = data_start + max(len(rows), 1) - 1 if rows else header_row
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(cols)}{last}"
    ws.freeze_panes = f"D{data_start}"


def _write_avg_vs_rec(ws: Any, payload: dict[str, Any]) -> None:
    headers = [
        "TIPO",
        "PRODUCTO + VARIANTE",
        "BARCODE",
        "STOCK",
        "COSTO PROM.",
        "ÚLT. COMP.",
        "Δ MONTO",
        "Δ VS PROM.",
        "BANDA",
        "FECHA",
        "N° FACTURA",
        "OFICINA",
        "COMPARABILIDAD",
        "EXTREMA",
        "PRIORIDAD",
        "ID VARIANTE",
        "SKU",
        "TIPO DOC.",
        "CONF. ERROR",
        "EVIDENCIA SUFICIENTE",
    ]
    rows = payload.get("average_vs_receptions") or []
    cols = len(headers)
    _set_widths(ws, [17, 51, 19, 11, 17, 16, 13, 13, 23, 14, 17, 18, 20, 12, 16, 13, 16, 14, 14, 18])
    _banner_pair(
        ws,
        title="COSTO PROMEDIO VS RECEPCIONES",
        subtitle=(
            f"{len(rows)} variantes en el alcance  ·  diferencia contra promedio Bsale, no entre facturas"
        ),
        cols=cols,
    )
    _merge_row(
        ws,
        3,
        cols,
        "TODAS LAS VARIANTES DEL ALCANCE  ·  Usa los filtros de «BANDA» y «COMPARABILIDAD». "
        "Sucursales con nombre administrativo (nunca ID).",
        fill=FILL_CHIP,
        font=MUTED_FONT,
    )
    header_row = 5
    _table_header(ws, headers, header_row)
    data_start = 6
    for offset, item in enumerate(rows):
        i = data_start + offset
        prio = item.get("admin_priority") or {}
        band = prio.get("review_priority") or prio.get("priority_band")
        diff = _num(item.get("diff_pct"))
        diff_ratio = (float(diff) / 100.0) if diff is not None else None
        diff_abs = _num(item.get("diff_abs"))
        values = [
            format_product_type(item.get("product_type_name")),
            format_product_plus_variant(item.get("product_name"), item.get("variant_name")),
            format_barcode(item.get("barcode") or item.get("sku")),
            _num(item.get("stock_total")),
            _num(item.get("average_cost")),
            _num(item.get("last_comparable_cost")),
            diff_abs if diff_abs is not None else None,
            diff_ratio,
            item.get("banda_principal"),
            _date(item.get("last_comparable_date")),
            _text(item.get("last_comparable_document_number")),
            _office_label(payload, item.get("last_comparable_office_id")),
            item.get("comparability_status"),
            "Sí" if item.get("extrema_indicador") else "No",
            _priority_label(band),
            item.get("variant_id"),
            _text(item.get("sku")),
            item.get("last_comparable_document"),
            prio.get("error_confidence"),
            "Sí" if item.get("evidence_sufficient") else "No",
        ]
        fill = None if item.get("evidence_sufficient") else FILL_PENDING
        zebra = fill is None and (offset % 2 == 1)
        for col, val in enumerate(values, start=1):
            cell = ws.cell(i, col, val)
            _style_data_cell(
                cell,
                money_dec=col in (5, 6, 7),
                pct_ratio=col == 8,
                as_date=col == 10,
                as_text=col in (3, 11, 17),
                fill=fill if col in (7, 8, 13, 20) else (_priority_fill(band) if col == 15 else None),
                zebra=zebra and fill is None and col != 15,
            )
            if col == 4 and cell.value is not None:
                cell.number_format = QTY_FMT
            if col == 15:
                pf = _priority_font(band)
                if pf:
                    cell.font = pf
        if not item.get("evidence_sufficient"):
            ws.cell(i, 7).value = None
            ws.cell(i, 8).value = None
            ws.cell(i, 7).fill = FILL_PENDING
            ws.cell(i, 8).fill = FILL_PENDING
        elif (
            diff_abs is None
            and _num(item.get("average_cost")) is not None
            and _num(item.get("last_comparable_cost")) is not None
        ):
            ws.cell(i, 7).value = f'=IF(OR(E{i}="",F{i}=""),"",F{i}-E{i})'
            _style_data_cell(ws.cell(i, 7), money_dec=True)

    last = data_start + max(len(rows), 1) - 1 if rows else header_row
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(cols)}{last}"
    ws.freeze_panes = f"D{data_start}"


def _write_integrity(ws: Any, payload: dict[str, Any]) -> None:
    _set_widths(ws, [47, 40, 50, 55])
    _banner_pair(
        ws,
        title="Integridad de catálogo y sincronización",
        subtitle="Cobertura de costos, huérfanas y frescura de sync — solo lectura",
        cols=4,
    )
    ws.cell(1, 1).font = Font(color="FFFFFF", bold=True, size=16)

    census = payload.get("census") or {}
    costs = payload.get("cost_coverage") or {}
    r = 3
    for idx, (label, value) in enumerate(
        [
            ("Variantes catálogo", census.get("variantes")),
            ("Filas variant_costs", costs.get("filas_variant_costs")),
            ("Sin fila de costo (aprox.)", (costs.get("variantes") or 0) - (costs.get("filas_variant_costs") or 0)),
            ("Costos NULL", costs.get("costs_average_null")),
            ("Costos = 0", costs.get("costs_average_zero")),
            ("Costos > 0", costs.get("costs_average_pos")),
            ("Activas sin fila stock", census.get("activas_sin_fila_stock")),
            ("Stock>0 sin variante (huérfanas)", census.get("stock_pos_sin_variante")),
            ("Unidades huérfanas", census.get("unidades_stock_pos_sin_variante")),
        ]
    ):
        zebra = idx % 2 == 1
        for col, val in enumerate((label, value), start=1):
            cell = ws.cell(r, col, val)
            _style_data_cell(cell, zebra=zebra)
        r += 1

    r += 1
    ws.cell(r, 1, "Huérfanas (detalle)").font = FONT_SECTION
    ws.cell(r, 1).fill = FILL_SECTION
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1
    _table_header(ws, ("ID VARIANTE", "STOCK", "OFICINAS", "NOTA"), r)
    r += 1
    orphans = payload.get("orphans") or []
    if not orphans:
        ws.cell(r, 1, "(sin huérfanas en esta fotografía)")
        r += 1
    for o in orphans:
        ws.cell(r, 1, o.get("variant_id"))
        ws.cell(r, 2, _num(o.get("stock_total")))
        ws.cell(r, 3, o.get("oficinas"))
        ws.cell(r, 4, "stock>0 sin fila en variants")
        r += 1

    r += 1
    ws.cell(r, 1, "Sincronización").font = FONT_SECTION
    ws.cell(r, 1).fill = FILL_SECTION
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1
    _table_header(ws, ("Recurso", "Estado", "Último SUCCESS", "Advertencia activa"), r)
    r += 1
    for s in payload.get("sync_state") or []:
        ws.cell(r, 1, f"{s.get('resource')} / {s.get('scope')}")
        ws.cell(r, 2, s.get("status"))
        ws.cell(r, 3, str(s.get("last_success_at") or ""))
        alert = "Sí" if s.get("active_freshness_alert") else "No"
        cell = ws.cell(r, 4, f"{alert} — {s.get('freshness_note')}")
        if s.get("active_freshness_alert"):
            cell.fill = FILL_P2
        r += 1

    r += 2
    ws.cell(r, 1, "Nota: errores históricos de sync no generan alerta si existe SUCCESS posterior.").font = FONT_MUTED
    ws.freeze_panes = "A3"


def _write_admin_review(ws: Any, payload: dict[str, Any]) -> None:
    headers = [
        "TIPO",
        "PRODUCTO + VARIANTE",
        "BARCODE",
        "PRIORIDAD",
        "MOTIVO",
        "FACTURA",
        "ESTADO",
        "RESPONSABLE",
        "FECHA REVISIÓN",
        "OBSERVACIÓN",
        "RESOLUCIÓN",
        "CORRECCIÓN APLICADA",
        "VARIANT_ID",
        "SKU",
    ]
    rows = payload.get("anomalies") or []
    cols = len(headers)
    cut = (payload.get("cut_at_utc") or "")[:10]
    _set_widths(ws, [17, 52, 19, 16, 25, 19, 23, 17, 14, 42, 32, 25, 13, 16])
    _banner_pair(
        ws,
        title="SEGUIMIENTO ADMINISTRATIVO",
        subtitle="Campos azules editables  ·  Cambiar estados en este archivo NO actualiza la ERP",
        cols=cols,
    )
    _merge_row(
        ws,
        3,
        cols,
        f"Una fila por variante. Registra responsable, fecha, decisión y evidencia. "
        f"Los casos vienen del corte del {cut or '—'}.",
        fill=FILL_PENDING,
        font=FONT_EDITABLE,
    )
    header_row = 5
    _table_header(ws, headers, header_row)
    data_start = 6
    for offset, item in enumerate(rows):
        i = data_start + offset
        prio = item.get("admin_priority") or {}
        band = prio.get("review_priority") or prio.get("priority_band")
        doc_n = item.get("last_comparable_document_number") or item.get("lc_document_number")
        doc_label = f"Factura {doc_n}" if doc_n else (
            item.get("last_comparable_document") or item.get("lc_document")
        )
        motivo = "Diferencia >15 %" + ("_extrema" if item.get("extrema_indicador") else "")
        values = [
            format_product_type(item.get("product_type_name")),
            format_product_plus_variant(item.get("product_name"), item.get("variant_name")),
            format_barcode(item.get("barcode") or item.get("sku")),
            _priority_label(band),
            motivo,
            doc_label,
            "Pendiente",
            None,
            None,
            None,
            None,
            None,
            item.get("variant_id"),
            _text(item.get("sku")),
        ]
        zebra = offset % 2 == 1
        for col, val in enumerate(values, start=1):
            cell = ws.cell(i, col, val)
            _style_data_cell(
                cell,
                as_text=col in (3, 14),
                as_date=col == 9,
                fill=_priority_fill(band) if col == 4 else (FILL_PENDING if col >= 7 and col <= 12 else None),
                zebra=zebra and col < 7,
            )
            if col == 4:
                pf = _priority_font(band)
                if pf:
                    cell.font = pf
            if 7 <= col <= 12:
                cell.font = FONT_EDITABLE

    if rows:
        last = data_start + len(rows) - 1
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(ADMIN_STATES) + '"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Estado inválido",
            error="Seleccione un estado de la lista.",
        )
        dv.add(f"G{data_start}:G{last}")
        ws.add_data_validation(dv)
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(cols)}{last}"
    ws.freeze_panes = f"D{data_start}"


def commercial_premium_xlsx(payload: dict[str, Any]) -> bytes:
    """Genera el workbook Premium de exactamente 6 hojas funcionales."""
    wb = Workbook()
    # Eliminar hoja default
    default = wb.active
    wb.remove(default)

    ws1 = wb.create_sheet(SHEET_NAMES[0])
    _write_summary(ws1, payload)

    ws2 = wb.create_sheet(SHEET_NAMES[1])
    _write_anomalies(ws2, payload)

    ws3 = wb.create_sheet(SHEET_NAMES[2])
    _write_zero_cost(ws3, payload)

    ws4 = wb.create_sheet(SHEET_NAMES[3])
    _write_avg_vs_rec(ws4, payload)

    ws5 = wb.create_sheet(SHEET_NAMES[4])
    _write_integrity(ws5, payload)

    ws6 = wb.create_sheet(SHEET_NAMES[5])
    _write_admin_review(ws6, payload)

    assert list(wb.sheetnames) == list(SHEET_NAMES)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _review_motivo(item: dict[str, Any]) -> str:
    prio = item.get("admin_priority") or {}
    action = prio.get("suggested_action") or item.get("recommended_action")
    source = item.get("review_source") or ""
    if source == "stock_positivo_sin_costo_valido":
        band = "stock positivo sin costo válido"
    elif source == "divergencia_y_stock_sin_costo":
        band = "divergencia >15 % · stock positivo sin costo válido"
    elif item.get("extrema_indicador"):
        band = "divergencia extrema"
    else:
        band = "divergencia >15 %"
    parts = [band]
    if action:
        parts.append(str(action))
    doc = " ".join(
        str(x)
        for x in (
            item.get("last_comparable_document") or item.get("lc_document"),
            item.get("last_comparable_document_number") or item.get("lc_document_number"),
        )
        if x
    ).strip()
    if doc:
        parts.append(f"Respaldo: {doc}")
    return " · ".join(parts)


def build_review_work_rows(payload: dict[str, Any]) -> list[dict[str, Any]]:
    """Filas accionables: anomalías >15 % + stock positivo sin costo válido. Una por variant_id."""
    by_vid: dict[Any, dict[str, Any]] = {}

    for a in payload.get("anomalies") or []:
        vid = a.get("variant_id")
        if vid is None:
            continue
        row = dict(a)
        row["review_source"] = "divergencia_gt_15"
        by_vid[vid] = row

    for z in payload.get("zero_cost_actives") or []:
        if z.get("stock_status") != "stock_positivo":
            continue
        vid = z.get("variant_id")
        if vid is None:
            continue
        if vid in by_vid:
            # Ya hay divergencia >15 %: no duplicar; marcar fuente mixta.
            existing = by_vid[vid]
            existing["review_source"] = "divergencia_y_stock_sin_costo"
            continue
        row = dict(z)
        row["review_source"] = "stock_positivo_sin_costo_valido"
        row["stock_total"] = z.get("stock_total_raw") if z.get("stock_total_raw") is not None else z.get("stock_total")
        row["average_cost"] = z.get("average_cost")
        row["diff_pct"] = None
        row["admin_priority"] = {
            "review_priority": "P1_critico",
            "priority_band": "P1_critico",
            "priority_score": 95,
            "suggested_action": z.get("recommended_action") or "revisar_costo_promedio_con_stock",
        }
        by_vid[vid] = row

    rows = list(by_vid.values())
    rows.sort(
        key=lambda r: (
            -float((r.get("admin_priority") or {}).get("priority_score") or 0),
            -abs(float(r.get("diff_pct") or 0)),
            int(r.get("variant_id") or 0),
        )
    )
    return rows


def _write_review_actionable(ws: Any, payload: dict[str, Any]) -> None:
    """Hoja operativa estilo PRO: banner, chips P1/P2/P3, tabla editable."""
    ws.sheet_view.showGridLines = False
    cols = 11
    headers = [
        "TIPO DE PRODUCTO",
        "PRODUCTO + VARIANTE",
        "CÓDIGO DE BARRAS",
        "STOCK GLOBAL",
        "COSTO PROM. NETO",
        "ÚLTIMA COMPRA NETA",
        "Δ VS PROMEDIO",
        "PRIORIDAD",
        "QUÉ REVISAR / FACTURA",
        "ESTADO",
        "OBSERVACIÓN",
    ]
    _set_widths(ws, [18, 52, 18, 11, 17, 18, 14, 16, 37, 23, 38])
    rows = build_review_work_rows(payload)
    n = len(rows)
    p1 = sum(
        1
        for a in rows
        if str((a.get("admin_priority") or {}).get("review_priority") or "").upper().startswith("P1")
        or a.get("review_source") == "stock_positivo_sin_costo_valido"
    )
    p2 = sum(
        1
        for a in rows
        if str((a.get("admin_priority") or {}).get("review_priority") or "").upper().startswith("P2")
    )
    p3 = max(0, n - p1 - p2)
    stock_label = payload.get("stock_filter_label") or STOCK_FILTER_LABELS.get(
        payload.get("stock_filter") or "with", "CON STOCK"
    )
    cut = (payload.get("cut_at_utc") or "")[:10]
    company = payload.get("company_name") or f"Empresa {payload.get('company_id')}"

    _banner_pair(
        ws,
        title="LA QUILLOTANA  /  CONTROL DE COSTOS",
        subtitle=(
            f"HOJA DE TRABAJO  ·  {n} casos para revisar  ·  {company}  ·  Corte {cut}"
        ),
        cols=cols,
    )
    # Chips de prioridad (fila 3)
    chips = [
        (1, 3, f"{n}  CASOS DEL CORTE", FILL_CHIP, INK_BOLD),
        (4, 5, f"{p1}  CRÍTICOS P1", FILL_P1, FONT_P1),
        (6, 7, f"{p2}  ALTOS P2", FILL_P2, INK_BOLD),
        (8, 9, f"{p3}  MEDIOS P3", FILL_P3, FONT_P3),
        (10, 11, f"{stock_label}  ✓", FILL_STOCK_OK, FONT_STOCK_OK),
    ]
    for c1, c2, text, fill, font in chips:
        ws.merge_cells(start_row=3, start_column=c1, end_row=3, end_column=c2)
        cell = ws.cell(3, c1, text)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", horizontal="center")
        for c in range(c1, c2 + 1):
            ws.cell(3, c).fill = fill
    ws.row_dimensions[3].height = 22

    _merge_row(
        ws,
        4,
        cols,
        "PASO 1  Filtra por prioridad   →   PASO 2  Revisa factura/unidad   →   PASO 3  Cambia ESTADO y escribe OBSERVACIÓN",
        fill=PatternFill("solid", fgColor="FFFFFF"),
        font=MUTED_FONT,
    )
    _merge_row(
        ws,
        5,
        cols,
        "Δ % = (última compra comparable − costo promedio Bsale) / costo promedio. No es variación entre dos facturas.",
        fill=FILL_WARN,
        font=INK_FONT,
    )
    _merge_row(
        ws,
        6,
        cols,
        "ALERTA ≠ ERROR CONFIRMADO   |   Los campos AZULES se pueden editar. Editar aquí NO actualiza la ERP.",
        fill=FILL_PENDING,
        font=FONT_EDITABLE,
    )
    # Fila 7 vacía separadora
    header_row = 8
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = THIN
    data_start = header_row + 1

    for offset, item in enumerate(rows):
        i = data_start + offset
        prio = item.get("admin_priority") or {}
        band = prio.get("review_priority") or prio.get("priority_band")
        fill = _priority_fill(band)
        zebra = (offset % 2 == 1)
        diff = _num(item.get("diff_pct"))
        diff_ratio = (float(diff) / 100.0) if diff is not None else None
        doc = item.get("last_comparable_document_number") or item.get("lc_document_number")
        motivo = _review_motivo(item)
        if doc:
            motivo = f"{motivo}\nFactura {doc}" if motivo else f"Factura {doc}"
        values: list[Any] = [
            format_product_type(item.get("product_type_name")),
            format_product_plus_variant(item.get("product_name"), item.get("variant_name")),
            format_barcode(item.get("barcode")),
            _num(item.get("stock_total")),
            _num(item.get("average_cost")),
            _num(item.get("last_comparable_cost")),
            diff_ratio,
            _priority_label(band),
            motivo,
            "Pendiente",
            None,
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(i, col, val)
            base_fill = FILL_PENDING if col >= 10 else (fill if col in (7, 8) else None)
            _style_data_cell(
                cell,
                as_text=col == 3,
                money=False,
                money_dec=col in (5, 6),
                pct_ratio=col == 7,
                fill=base_fill,
                zebra=zebra and base_fill is None,
            )
            if col == 4 and cell.value is not None:
                cell.number_format = QTY_FMT
                cell.font = INK_BOLD
            if col == 8:
                pf = _priority_font(band)
                if pf:
                    cell.font = pf
            if col == 1:
                cell.font = Font(color="536276", bold=True, size=9)
            if col == 2:
                cell.font = INK_BOLD
            if col >= 10:
                cell.font = FONT_EDITABLE

    last_data = data_start + max(n, 1) - 1 if n else data_start
    if n:
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(ADMIN_STATES) + '"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Estado inválido",
            error="Seleccione un estado de la lista.",
        )
        dv.add(f"J{data_start}:J{last_data}")
        ws.add_data_validation(dv)
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(cols)}{last_data}"
    # Identidad A–C + bloque superior (banner/chips/instrucciones/header) fijos.
    ws.freeze_panes = f"D{data_start}"


def _write_review_summary(ws: Any, payload: dict[str, Any]) -> None:
    """Resumen operativo estilo PRO con progreso COUNTIF sobre «Por revisar»."""
    ws.sheet_view.showGridLines = False
    _set_widths(ws, [14, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14, 15])
    work = build_review_work_rows(payload)
    n = len(work)
    census = payload.get("census") or {}
    buckets = payload.get("divergence_buckets") or {}
    anomalies = payload.get("anomalies") or []
    zc = payload.get("zero_cost_summary") or {}
    stock_sin_costo = int(zc.get("stock_positivo") or 0)
    critical = sum(
        1
        for a in work
        if str((a.get("admin_priority") or {}).get("review_priority") or "").upper().startswith("P1")
        or a.get("extrema_indicador")
        or a.get("review_source") == "stock_positivo_sin_costo_valido"
    )
    cut = (payload.get("cut_at_utc") or "")[:10]
    company_id = payload.get("company_id")
    stock_pos = census.get("activas_stock_pos")
    data_start = 9
    data_end = data_start + max(n, 1) - 1

    _banner_pair(
        ws,
        title="LA QUILLOTANA  /  REVISIÓN INTELIGENTE",
        subtitle=f"INFORME OPERATIVO DE COSTOS  ·  EMPRESA {company_id}  ·  FOTO AL {cut}",
        cols=12,
    )

    # KPI cards
    def _kpi(r_val: int, r_lab: int, c: int, value: Any, label: str, fill: PatternFill, font: Font) -> None:
        ws.merge_cells(start_row=r_val, start_column=c, end_row=r_val, end_column=c + 3)
        ws.merge_cells(start_row=r_lab, start_column=c, end_row=r_lab, end_column=c + 3)
        v = ws.cell(r_val, c, value)
        v.fill = fill
        v.font = Font(color=font.color.rgb if font.color and font.color.rgb else "1B2637", bold=True, size=24)
        v.alignment = Alignment(horizontal="center", vertical="center")
        lab = ws.cell(r_lab, c, label)
        lab.fill = fill
        lab.font = font
        lab.alignment = Alignment(horizontal="center", vertical="center")
        for cc in range(c, c + 4):
            ws.cell(r_val, cc).fill = fill
            ws.cell(r_lab, cc).fill = fill

    _kpi(4, 5, 1, n, "ALERTAS POR REVISAR", FILL_CHIP, INK_BOLD)
    _kpi(4, 5, 5, critical, "P1 CRÍTICOS", FILL_P1, FONT_P1)
    _kpi(4, 5, 9, stock_pos, "VARIANTES CON STOCK¹", FILL_STOCK_OK, FONT_STOCK_OK)
    _kpi(7, 8, 1, critical, "CRÍTICOS² (P1 O EXTREMOS)", FILL_P1, FONT_P1)
    _kpi(
        7,
        8,
        5,
        buckets.get("requiere_verificacion"),
        "PENDIENTES DE VERIFICAR¹",
        FILL_WARN,
        INK_BOLD,
    )
    _kpi(
        7,
        8,
        9,
        buckets.get("sin_recepcion_comparable"),
        "SIN COMPRA COMPARABLE¹",
        FILL_P3,
        FONT_P3,
    )

    _merge_row(
        ws,
        10,
        12,
        "PROGRESO DE REVISIÓN — se actualiza al cambiar los estados de la hoja «Por revisar»",
        fill=BANNER_SUB_FILL,
        font=FONT_SECTION,
    )
    for c, label in ((1, "PENDIENTES"), (5, "EN REVISIÓN"), (9, "CERRADOS")):
        ws.merge_cells(start_row=11, start_column=c, end_row=11, end_column=c + 3)
        cell = ws.cell(11, c, label)
        cell.fill = FILL_CHIP
        cell.font = MUTED_FONT
        cell.alignment = Alignment(horizontal="center")
        for cc in range(c, c + 4):
            ws.cell(11, cc).fill = FILL_CHIP

    sheet = REVIEW_SHEET_NAMES[0]
    formulas = [
        (1, f"=COUNTIF('{sheet}'!$J${data_start}:$J${data_end},\"Pendiente\")"),
        (5, f"=COUNTIF('{sheet}'!$J${data_start}:$J${data_end},\"En revisión\")"),
        (
            9,
            f"=COUNTIF('{sheet}'!$J${data_start}:$J${data_end},\"Revisado sin error\")"
            f"+COUNTIF('{sheet}'!$J${data_start}:$J${data_end},\"Corrección aplicada\")",
        ),
    ]
    for c, formula in formulas:
        ws.merge_cells(start_row=12, start_column=c, end_row=12, end_column=c + 3)
        cell = ws.cell(12, c, formula)
        cell.font = Font(color="1B2637", bold=True, size=21)
        cell.alignment = Alignment(horizontal="center")

    _merge_row(
        ws,
        14,
        12,
        "CÓMO USAR ESTE ARCHIVO  /  EN 3 PASOS",
        fill=BANNER_SUB_FILL,
        font=FONT_SECTION,
    )
    steps = [
        "1. En «Por revisar», filtre por PRIORIDAD (P1 primero).",
        "2. Contraste factura / unidad / presentación (SEC) con el costo promedio Bsale.",
        "3. Cambie ESTADO y deje OBSERVACIÓN. Editar el Excel NO actualiza la ERP ni Bsale.",
        f"Diferencias >15 % en este corte: {payload.get('anomalies_count', len(anomalies))} · "
        f"Stock positivo sin costo válido: {stock_sin_costo}.",
        "¹ Universo operativo CON STOCK (no mezclar con totales GLOBAL de catálogo).",
        "² Críticos = prioridad P1, extremos o stock positivo sin costo válido.",
        "Las alertas son señales de revisión; no son errores confirmados ni órdenes automáticas.",
        "ALERTA ≠ ERROR CONFIRMADO. Editar el Excel no constituye corrección en Bsale ni PostgreSQL.",
    ]
    r = 15
    for text in steps:
        _merge_row(ws, r, 12, text, fill=PatternFill("solid", fgColor="FFFFFF"), font=INK_FONT)
        r += 1
    ws.freeze_panes = "A13"


def commercial_premium_review_xlsx(payload: dict[str, Any]) -> bytes:
    """Workbook operativo de 2 hojas (revisión diaria). Reutiliza el mismo payload."""
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    ws1 = wb.create_sheet(REVIEW_SHEET_NAMES[0])
    _write_review_actionable(ws1, payload)

    ws2 = wb.create_sheet(REVIEW_SHEET_NAMES[1])
    _write_review_summary(ws2, payload)

    assert list(wb.sheetnames) == list(REVIEW_SHEET_NAMES)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
